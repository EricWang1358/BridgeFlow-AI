"""Immutable uploaded batches. Row-bearing views are browser-only, never tools."""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Annotated

import openpyxl
import pandas as pd
import yaml
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field

from bridgeflow import column_matches
from bridgeflow.access import departments_for, operations_for
from bridgeflow.agents import DataSanitizerAgent, SanitizerInput
from bridgeflow.agents.semantic_resolver import FieldDictionary, SemanticResolverAgent
from bridgeflow.agents.sop_flow import MissingRollup, SOPFlowEngine, SOPInput, UnjoinableTables
from bridgeflow.column_matches import AppliedMatch
from bridgeflow.config import REPO_ROOT, settings
from bridgeflow.identity import UserIdentity, require_user
from bridgeflow.metrics import dictionary_path
from bridgeflow.monthly import checks
from bridgeflow.schemas import CleanTable, Department, PipelineResult
from bridgeflow.store import _root, _write

router = APIRouter(prefix="/batches", tags=["batches"])


def _json_cell(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value.item() if hasattr(value, "item") else value


def batch_path(batch_id: str):
    if not re.fullmatch(r"[a-f0-9]{32}", batch_id):
        raise HTTPException(422, "Invalid batch id")
    return _root() / "batches" / f"{batch_id}.json"


class BatchSnapshot(PipelineResult):
    dictionary_snapshot: dict | None = None
    integration_snapshot: dict | None = None
    refusal: str = ""
    demo_case: str | None = None
    dictionary_source: str | None = None
    #: Uploaded columns renamed onto declared columns at import, from human decisions.
    column_matches: list[AppliedMatch] = Field(default_factory=list)
    #: Decisions not applied because the uploaded column's shape changed.
    stale_matches: list[str] = Field(default_factory=list)
    #: The frozen batch this one was derived from by applying quarantine decisions (#88).
    derived_from: str | None = None
    dispositions: list[dict] = Field(default_factory=list)
    #: Per department, the shared intake check report (E14-UC03): identical to what self-check returns.
    intake_checks: dict[str, dict] = Field(default_factory=dict)
    #: Portal union_id of whoever imported it; "" = system/tool import or predates
    #: the identity layer. Visibility: the batch's departments ⊆ the viewer's
    #: authorized departments, or the viewer is the owner (docs/27).
    owner: str = ""
    #: Intake columns dropped because nothing in them was representable (docs/33, FR-5).
    #: Column-level fact, kept apart from quarantine, which is row-level disposition.
    dropped_columns: list[dict] = Field(default_factory=list)


def _visible(batch: BatchSnapshot, user: UserIdentity | None) -> BatchSnapshot:
    """The batch if this identity may see it. Invisible means 404, never 403:
    whether a batch exists is itself information."""
    if user is None or batch.owner == user.sub:
        return batch
    if {table.department for table in batch.clean_tables} <= departments_for(user.sub):
        return batch
    raise HTTPException(404, "Batch not found")


def load_batch(batch_id: str) -> BatchSnapshot:
    path = batch_path(batch_id)
    if not path.is_file():
        raise HTTPException(404, "Batch not found")
    return BatchSnapshot.model_validate_json(path.read_text(encoding="utf-8"))


def batch_dictionary(batch: BatchSnapshot) -> FieldDictionary:
    if batch.dictionary_snapshot is None:
        raise HTTPException(409, "Batch has no dictionary snapshot; import a new batch")
    return FieldDictionary(batch.dictionary_snapshot)


class BatchRef(BaseModel):
    batch_id: str


class DepartmentSummary(BaseModel):
    department: str
    rows: int
    quarantined: int
    corrections: int


class BatchSummary(BaseModel):
    demo_case: str | None = None
    batch_id: str
    period: str
    departments: list[DepartmentSummary]
    master_rows: int
    unresolved: int
    status: str
    refusal: str = ""
    #: The dictionary this batch was frozen against, and what it declares as a
    #: joinable entity per department.
    #:
    #: A refusal names the department that has no joinable column, which is the right
    #: half of the answer and useless on its own: somebody staring at
    #: `needs_configuration` cannot tell whether the dictionary is wrong or whether
    #: they are simply pointed at a different dictionary than they think. Measured on
    #: `data/business_demo/risk`: its finance sheet declares `project`, the default
    #: dictionary declares `gl_account`, and nothing on screen said which file was in
    #: force — so the walkthrough's step 1 passed and step 2 was impossible.
    dictionary: str = ""
    declared_entities: dict[str, list[str]] = Field(default_factory=dict)
    #: `department.column → declared column`, applied from remembered decisions.
    matched_columns: list[str] = Field(default_factory=list)
    stale_matches: list[str] = Field(default_factory=list)
    #: Uploaded columns the dictionary does not know that might be a declared column
    #: this department is missing. Non-zero means the captain has something to propose.
    column_questions: int = 0
    derived_from: str | None = None
    #: Intake columns dropped as unrepresentable (docs/33); empty for plain file uploads.
    dropped_columns: list[dict] = Field(default_factory=list)


def _declared_entities(snapshot: dict | None) -> dict[str, list[str]]:
    """Which columns each department declares as an entity, straight from the snapshot.

    Column names only — never a value out of anybody's sheet. This is what the
    dictionary says, not what the data holds.
    """
    columns = (snapshot or {}).get("columns") or {}
    if not isinstance(columns, dict):
        return {}
    return {
        str(department): sorted(str(name) for name in declared)
        for department, declared in columns.items()
        if isinstance(declared, dict) and declared
    }


def summary(batch_id: str, result: BatchSnapshot) -> BatchSummary:
    return BatchSummary(
        batch_id=batch_id, period=result.period,
        departments=[DepartmentSummary(
            department=t.department, rows=len(t.rows), quarantined=len(t.quarantine),
            corrections=len(t.corrections),
        ) for t in result.clean_tables],
        master_rows=len(result.master_table.rows) if result.master_table else 0,
        unresolved=len(result.graph.unresolved) if result.graph else 0,
        status=("needs_configuration" if result.refusal else "needs_review"
                if result.graph.unresolved or any(t.quarantine for t in result.clean_tables)
                else "ready" if result.master_table.rows else "empty"),
        refusal=result.refusal,
        demo_case=result.demo_case,
        dictionary=result.dictionary_source or _relative_dictionary(),
        declared_entities=_declared_entities(result.dictionary_snapshot),
        matched_columns=[f"{m.department}.{m.column} → {m.target}" for m in result.column_matches],
        stale_matches=result.stale_matches,
        column_questions=column_matches.count_questions(result.dictionary_snapshot, result.clean_tables),
        derived_from=result.derived_from,
        dropped_columns=result.dropped_columns,
    )


def _relative_dictionary() -> str:
    """The dictionary path as somebody would type it, not as the process sees it."""
    path = dictionary_path()
    try:
        return str(path.relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def _check_upload_scope(departments: list[Department], user: UserIdentity | None) -> None:
    """With the identity layer on, you may only upload departments you may also see."""
    if user is None:
        return
    if "batch_import" not in operations_for(user.sub):
        raise HTTPException(403, "This employee is not authorized to import batches")
    denied = set(departments) - departments_for(user.sub)
    if denied:
        raise HTTPException(403, f"Not authorized for departments: {sorted(denied)}")


@router.post("", response_model=BatchSummary)
async def upload_batch(
    period: Annotated[str, Form(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")],
    departments: Annotated[list[Department], Form()],
    files: Annotated[list[UploadFile], File()],
    user: Annotated[UserIdentity | None, Depends(require_user)],
    sheets: Annotated[list[str] | None, Form(description="Per file, the worksheet to read; empty uses the declaration")] = None,
    header_rows: Annotated[list[str] | None, Form(description="Per file, the 1-based header row; empty uses the declaration")] = None,
) -> BatchSummary:
    _check_upload_scope(departments, user)
    choices = [Layout.from_form(sheet, header) for sheet, header in
               zip(_aligned(sheets, len(files)), _aligned(header_rows, len(files)), strict=True)]
    return await _import_batch(period, departments, files, choices=choices,
                               owner=user.sub if user else "")


def _aligned(values: list[str] | None, count: int) -> list[str]:
    if not values:
        return [""] * count
    if len(values) != count:
        raise HTTPException(422, "sheets and header_rows need one entry per file (empty for none)")
    return values


@dataclass
class Layout:
    """Which worksheet holds the table and which row is its header — declared or chosen, never guessed."""
    sheet: str = ""
    header_row: int | None = None

    @classmethod
    def from_form(cls, sheet: str, header: str) -> Layout:
        if header.strip() and not (header.strip().isdigit() and int(header) >= 1):
            raise HTTPException(422, "header_rows entries must be positive row numbers")
        return cls(sheet.strip(), int(header) if header.strip() else None)

    @classmethod
    def declared(cls, raw: object, department: str) -> Layout:
        entry = raw.get(department) if isinstance(raw, dict) else None
        if entry is None:
            return cls()
        if not isinstance(entry, dict) or set(entry) - {"sheet", "header_row"} \
                or not isinstance(entry.get("sheet", ""), str) \
                or not (entry.get("header_row") is None or (isinstance(entry["header_row"], int) and entry["header_row"] >= 1)):
            raise ValueError("sheet_layout entries declare sheet (text) and header_row (positive integer)")
        return cls(entry.get("sheet", ""), entry.get("header_row"))

    def over(self, declared: Layout) -> Layout:
        return Layout(self.sheet or declared.sheet, self.header_row or declared.header_row)


def _read_xlsx(payload: bytes, filename: str, layout: Layout) -> tuple[str, pd.DataFrame, int]:
    """One worksheet as a table. Ambiguity is refused with what a person needs to choose."""
    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    names = workbook.sheetnames
    if layout.sheet:
        if layout.sheet not in names:
            raise HTTPException(422, f"{filename}: no worksheet named {layout.sheet!r}; it has {names}")
        sheet = layout.sheet
    elif len(names) == 1:
        sheet = names[0]
    else:
        raise HTTPException(422, f"{filename}: multiple sheets {names}; choose one per file (sheets) or declare sheet_layout")
    # A formula written by a program and never opened in a spreadsheet app has no saved
    # result; reading it would silently turn every such cell into a blank.
    formulas = workbook[sheet].iter_rows(values_only=True)
    cached = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    top: list[int] = []
    unsaved: list[str] = []
    for index, (written, stored) in enumerate(zip(formulas, cached[sheet].iter_rows(values_only=True), strict=False)):
        if index < 10:
            top.append(sum(c not in (None, "") for c in stored))
        if any(isinstance(f, str) and f.startswith("=") and c is None for f, c in zip(written, stored, strict=False)):
            unsaved.append(f"row {index + 1}")
            if len(unsaved) == 3:
                break
    workbook.close()
    cached.close()
    if unsaved:
        raise HTTPException(422, f"{filename}: formulas without saved results ({', '.join(unsaved)}); "
                                 "open and save the file in Excel or WPS so their values are stored, then upload again")
    header_row = layout.header_row
    if header_row is None:
        # Row 1 is the header only when it looks like one; a title or blank row above the
        # table is not resolved by guessing which later row is meant.
        if top and top[0] <= 1 and any(n > 1 for n in top[1:]):
            candidate = next(i for i, n in enumerate(top) if n > 1) + 1
            raise HTTPException(422, f"{filename}: row 1 is not a header (row {candidate} looks like one); "
                                     "confirm the header row per file (header_rows) or declare sheet_layout")
        header_row = 1
    frame = pd.read_excel(io.BytesIO(payload), sheet_name=sheet, header=header_row - 1,
                          nrows=settings.bridgeflow_max_batch_rows + 1)
    return sheet, frame, header_row


@router.post("/demo", response_model=BatchSummary)
async def demo_batch(user: Annotated[UserIdentity | None, Depends(require_user)]) -> BatchSummary:
    """Explicit sample notebook; freeze its dictionary without replacing deployment policy.

    The sample is the fictional concrete supplier in `data/mock_business/demo`, filed on the
    business side's v2 department templates, so the notebook shows the templates, the
    cross-department master table and the declared review together.
    """
    folder = REPO_ROOT / "data/mock_business/demo"
    departments: list[Department] = ["production", "procurement", "finance", "marketing"]
    labels = {"production": "生产部", "procurement": "物资部", "finance": "财务部", "marketing": "市场部"}
    _check_upload_scope(departments, user)
    files = [UploadFile(io.BytesIO((folder / f"{department}.xlsx").read_bytes()),
                        filename=f"模拟-{labels[department]}-2024-07.xlsx") for department in departments]
    try:
        return await _import_batch("2024-07", departments, files, folder / "dictionary.yaml",
                                   "mock-company-2024-07", owner=user.sub if user else "")
    finally:
        for upload in files:
            await upload.close()


def _load_dictionary(path: Path) -> tuple[dict, FieldDictionary, dict]:
    try:
        dictionary_raw = yaml.safe_load(path.read_text(encoding="utf-8")) if path.is_file() else {}
        if dictionary_raw is None:
            dictionary_raw = {}
        if not isinstance(dictionary_raw, dict):
            raise TypeError("Dictionary root must be a mapping")
        dictionary = FieldDictionary(dictionary_raw)
        # How each department writes a slashed date is the dictionary owner's declaration (#79).
        date_orders = dictionary_raw.get("date_order") or {}
        if not isinstance(date_orders, dict) or not set(date_orders.values()) <= {"day_first", "month_first"}:
            raise ValueError("date_order must map departments to day_first or month_first")
    except (yaml.YAMLError, OSError, UnicodeError, ValueError, TypeError, AttributeError, KeyError) as exc:
        raise HTTPException(503, "Field dictionary configuration is invalid; ask the administrator to correct it before importing") from exc
    return dictionary_raw, dictionary, date_orders


def _integration_spec():
    # Imported locally because the evaluator also uses BatchSnapshot for its arithmetic.
    from bridgeflow import integration

    return integration.load_spec() if integration.spec_path().is_file() else None


def _declaration_label(dictionary_raw: dict, spec) -> str:
    digest = hashlib.sha256(json.dumps(dictionary_raw, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()[:12]
    return f"dictionary {digest}" + (f" · {spec.version}" if spec is not None else "")


@dataclass
class ParsedFile:
    department: str
    filename: str
    sheet: str
    header_row: int
    frame: pd.DataFrame
    sha256: str
    size: int


def _parse_department_file(department: str, filename: str, payload: bytes, layout: Layout) -> ParsedFile:
    """Read one department file exactly as import does; refusals are HTTP errors with a reason."""
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
    sheet = ""
    header_row = 1
    try:
        if filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise HTTPException(413, "Expanded workbook exceeds 100 MiB")
            sheet, frame, header_row = _read_xlsx(payload, filename, layout)
        elif filename.lower().endswith(".csv"):
            frame = pd.read_csv(io.BytesIO(payload), skip_blank_lines=False, nrows=settings.bridgeflow_max_batch_rows + 1)
        else:
            raise HTTPException(415, "Use CSV or an XLSX workbook")
    except HTTPException:
        raise
    except (ValueError, OSError, zipfile.BadZipFile, UnicodeError) as exc:
        raise HTTPException(422, f"Cannot parse {filename}") from exc
    if len(frame) > settings.bridgeflow_max_batch_rows:
        raise HTTPException(413, "Batch exceeds configured row limit")
    if frame.empty:
        raise HTTPException(422, f"{filename} has no data rows")
    return ParsedFile(department, filename, sheet, header_row, frame, hashlib.sha256(payload).hexdigest(), len(payload))


def _source_record(parsed: ParsedFile) -> dict:
    # Browser-only originals live outside the model-readable batch snapshot.
    # Capture before sanitation mutates values; this is a parsed table preview.
    # Python's own JSON keeps every double exactly; pandas' writer caps at 15 digits, and
    # the retained original must not be a rounded copy of the sheet.
    frame = parsed.frame
    return {"id": parsed.department, "filename": parsed.filename, "sheet": parsed.sheet,
            "sha256": parsed.sha256, "bytes": parsed.size,
            "columns": [str(c) for c in frame.columns],
            "rows": [[_json_cell(v) for v in row] for row in frame.astype(object).where(frame.notna(), None).values.tolist()],
            "header_row": parsed.header_row,
            "row_numbers": list(range(parsed.header_row + 1, parsed.header_row + 1 + len(frame)))}


async def _clean(parsed: ParsedFile, period: str, date_order: str | None, batch_id: str) -> CleanTable:
    table = await DataSanitizerAgent().run(SanitizerInput(parsed.department, period, parsed.frame, date_order))
    table.filename, table.sheet, table.batch = parsed.filename, parsed.sheet, batch_id
    table.source_rows = [n + parsed.header_row - 1 for n in table.source_rows]
    for correction in table.corrections:
        correction.source.filename = parsed.filename
        correction.source.sheet = parsed.sheet
        correction.source.batch = batch_id
        correction.source.source_row = correction.row + parsed.header_row + 1
    return table


@router.post("/self-check", response_model=checks.CheckReport)
async def self_check(
    period: Annotated[str, Form(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")],
    department: Annotated[Department, Form()],
    file: Annotated[UploadFile, File()],
    user: Annotated[UserIdentity | None, Depends(require_user)],
    sheet: Annotated[str, Form()] = "",
    header_row: Annotated[str, Form()] = "",
) -> checks.CheckReport:
    """Run import's own check chain on one department file without creating a batch (E14-UC03).

    Nothing is written: no batch, no retained original, no mapping memory.
    """
    _check_upload_scope([department], user)
    dictionary_raw, _, date_orders = _load_dictionary(dictionary_path())
    try:
        declared_layout = Layout.declared(dictionary_raw.get("sheet_layout"), department)
    except (ValueError, TypeError, AttributeError, KeyError) as exc:
        raise HTTPException(503, "Field dictionary configuration is invalid; ask the administrator to correct it before importing") from exc
    spec = _integration_spec()
    declared = _declaration_label(dictionary_raw, spec)
    payload = await file.read(settings.bridgeflow_max_upload_bytes + 1)
    filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    if len(payload) > settings.bridgeflow_max_upload_bytes:
        return checks.CHAIN.run(checks.CheckSubject(department, filename, declared, refusal="Batch exceeds configured upload size limit"))
    try:
        parsed = _parse_department_file(department, filename, payload, Layout.from_form(sheet, header_row).over(declared_layout))
    except HTTPException as exc:
        return checks.CHAIN.run(checks.CheckSubject(department, filename, declared, refusal=str(exc.detail)))
    table = await _clean(parsed, period, date_orders.get(department), "")
    return checks.CHAIN.run(checks.CheckSubject(department, parsed.filename, declared, table=table,
                                                source=_source_record(parsed), spec=spec))


async def _import_batch(period: str, departments: list[Department], files: list[UploadFile],
                        source_dictionary: Path | None = None, demo_case: str | None = None,
                        choices: list[Layout] | None = None, owner: str = "",
                        dropped_columns: list[dict] | None = None) -> BatchSummary:
    if len(files) != len(departments) or not 1 <= len(files) <= 4:
        raise HTTPException(422, "Provide one file per department (1–4 departments)")
    if len(set(departments)) != len(departments):
        raise HTTPException(422, "Duplicate department; combine its sheets explicitly first")
    batch_id = uuid.uuid4().hex
    path = source_dictionary or dictionary_path()
    dictionary_raw, dictionary, date_orders = _load_dictionary(path)
    try:
        # Where each department's table sits in its workbook, declared by the dictionary owner (#47).
        layouts = [Layout.declared(dictionary_raw.get("sheet_layout"), d) for d in departments]
    except (ValueError, TypeError, AttributeError, KeyError) as exc:
        raise HTTPException(503, "Field dictionary configuration is invalid; ask the administrator to correct it before importing") from exc
    spec = _integration_spec()
    integration_snapshot = spec.model_dump(mode="json") if spec is not None else None
    declared = _declaration_label(dictionary_raw, spec)
    tables = []
    sources = []
    byte_count = row_count = 0
    intake_reports: dict[str, dict] = {}
    for department, upload, declared_layout, chosen in zip(departments, files, layouts, choices or [Layout()] * len(files), strict=True):
        payload = await upload.read(settings.bridgeflow_max_upload_bytes + 1)
        byte_count += len(payload)
        if byte_count > settings.bridgeflow_max_upload_bytes:
            raise HTTPException(413, "Batch exceeds configured upload size limit")
        parsed = _parse_department_file(department, upload.filename or "", payload, chosen.over(declared_layout))
        frame = parsed.frame
        row_count += len(frame)
        if row_count > settings.bridgeflow_max_batch_rows:
            raise HTTPException(413, "Batch exceeds configured row limit")
        source = _source_record(parsed)
        sources.append(source)
        table = await _clean(parsed, period, date_orders.get(department), batch_id)
        report = checks.CHAIN.run(checks.CheckSubject(department, parsed.filename, declared, table=table,
                                                      source=source, spec=spec))
        intake_reports[department] = report.model_dump(mode="json")
        tables.append(table)
    result = BatchSnapshot(period=period, clean_tables=tables, dictionary_snapshot=dictionary_raw or {},
                           integration_snapshot=integration_snapshot, owner=owner, intake_checks=intake_reports,
                           dropped_columns=dropped_columns or [])
    if not dictionary.is_empty:
        # Before anything reads the tables: a remembered match makes this upload look
        # the way the dictionary already describes it. Frozen batches are never redone.
        result.column_matches, result.stale_matches = column_matches.apply(tables, dictionary_raw)
    await assemble(result, dictionary)
    result.demo_case = demo_case
    result.dictionary_source = str(path.relative_to(REPO_ROOT)) if path.is_relative_to(REPO_ROOT) else str(path)
    manifest = [{key: source[key] for key in ("id", "filename", "sheet", "sha256", "bytes")}
                | {"total": len(source["rows"])} for source in sources]
    _write(batch_path(batch_id).parent / "sources" / batch_id / "index.json", {"sources": manifest})
    for source in sources:
        _write(batch_path(batch_id).parent / "sources" / batch_id / f"{source['id']}.json", source)
    _write(batch_path(batch_id), result.model_dump(mode="json"))
    return summary(batch_id, result)


async def assemble(result: BatchSnapshot, dictionary: FieldDictionary) -> None:
    """Resolve entities and build the master table from cleaned tables, or record why not."""
    result.refusal = ""
    if dictionary.is_empty:
        result.refusal = "Field dictionary is not configured; declare fields and import a new batch"
        return
    result.graph = await SemanticResolverAgent().run(result.clean_tables, adjudicate=False, dictionary=dictionary)
    try:
        assembled = await SOPFlowEngine().run(
            SOPInput(period=result.period, tables=result.clean_tables, graph=result.graph), dictionary=dictionary,
        )
        result.master_table = assembled.master_table
    except (UnjoinableTables, MissingRollup) as exc:
        result.refusal = str(exc)


@router.get("/{batch_id}", response_model=BatchSummary)
async def get_summary(batch_id: str,
                      user: Annotated[UserIdentity | None, Depends(require_user)]) -> BatchSummary:
    return summary(batch_id, _visible(load_batch(batch_id), user))


@router.get("/{batch_id}/view")
async def view(
    batch_id: str, user: Annotated[UserIdentity | None, Depends(require_user)],
    section: str = "master", offset: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=100),
) -> dict:
    result = _visible(load_batch(batch_id), user)
    if section == "master":
        rows = result.master_table.rows if result.master_table else []
    elif section == "mappings":
        rows = [x.model_dump() for x in result.graph.unresolved] if result.graph else []
    elif section == "columns":
        # The column-match wizard's view (#46): each unknown uploaded column against the
        # declared columns it could be, with the evidence and any decision on file.
        listed = column_matches.candidates(batch_id, result.period, result.clean_tables, result.dictionary_snapshot,
                                           applied=result.column_matches)
        rows = [{"department": q.department, "column": q.column, "original": q.original, "candidate": c.target,
                 "role": c.role, "type_fits": c.type_fits,
                 "shared_values": None if c.entity_overlap is None else f"{round(c.entity_overlap * 100)}% · {c.overlap_with}",
                 "decision": c.decided or "open"}
                for q in listed.questions for c in q.candidates]
        rows += [{"department": d, "column": "", "original": "", "candidate": "", "role": "", "type_fits": None,
                  "shared_values": None, "decision": "needs_dictionary_owner"} for d in listed.needs_dictionary_owner]
    elif section == "corrections":
        rows = [c.model_dump() for t in result.clean_tables for c in t.corrections]
    elif section == "quarantine":
        from bridgeflow import quarantine as dispositions
        decided = dispositions.decisions_by_row(batch_id)
        rows = [{"department": t.department, "index": index, "values": row,
                 "decision": decided.get((t.department, index), {}).get("action", "")}
                for t in result.clean_tables for index, row in enumerate(t.quarantine)]
    else:
        raise HTTPException(422, "Unknown batch view")
    return {"batch_id": batch_id, "total": len(rows), "offset": offset,
            "rows": rows[offset:offset + limit]}


@router.get("/{batch_id}/review")
async def review(batch_id: str, user: Annotated[UserIdentity | None, Depends(require_user)],
                 report_id: str | None = None) -> dict:
    from bridgeflow.api.reviews import saved_review
    _visible(load_batch(batch_id), user)
    return saved_review(batch_id, report_id)


@router.get("/{batch_id}/review-notes/{report_id}")
async def review_notes(batch_id: str, report_id: str,
                       user: Annotated[UserIdentity | None, Depends(require_user)]) -> dict:
    _visible(load_batch(batch_id), user)
    from bridgeflow.api.reviews import review_notes as notes
    return notes(batch_id, report_id)


@router.get("/{batch_id}/sources")
async def list_sources(batch_id: str,
                       user: Annotated[UserIdentity | None, Depends(require_user)]) -> dict:
    batch = _visible(load_batch(batch_id), user)
    sources = []
    folder = batch_path(batch_id).parent / "sources" / batch_id
    manifest_path = folder / "index.json"
    manifest = {item["id"]: item for item in json.loads(manifest_path.read_text(encoding="utf-8"))["sources"]} if manifest_path.is_file() else {}
    for table in batch.clean_tables:
        path = batch_path(batch_id).parent / "sources" / batch_id / f"{table.department}.json"
        if path.is_file() and table.department in manifest:
            sources.append(manifest[table.department] | {"preview_available": True})
        else:
            # Older batches have no retained original. Never substitute cleaned rows.
            sources.append({"id": table.department, "filename": table.filename,
                            "sheet": table.sheet, "preview_available": False})
    return {"batch_id": batch_id, "sources": sources}


@router.get("/{batch_id}/sources/{source_id}")
async def source_preview(batch_id: str, source_id: Department,
                         user: Annotated[UserIdentity | None, Depends(require_user)],
                         offset: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=100)) -> dict:
    batch = _visible(load_batch(batch_id), user)
    if source_id not in {table.department for table in batch.clean_tables}:
        raise HTTPException(404, "Source not found in this batch")
    path = batch_path(batch_id).parent / "sources" / batch_id / f"{source_id}.json"
    if not path.is_file():
        raise HTTPException(404, "Original preview was not retained for this batch; import a new batch")
    source = json.loads(path.read_text(encoding="utf-8"))
    rows = source.pop("rows")
    numbers = source.pop("row_numbers", list(range(2, len(rows) + 2)))
    return {**source, "batch_id": batch_id, "total": len(rows), "offset": offset,
            "rows": rows[offset:offset + limit], "row_numbers": numbers[offset:offset + limit]}


@router.get("/{batch_id}/artifacts")
async def list_artifacts(batch_id: str, user: Annotated[UserIdentity | None, Depends(require_user)],
                         offset: int = Query(0, ge=0),
                         limit: int = Query(50, ge=1, le=100)) -> dict:
    _visible(load_batch(batch_id), user)
    folder = batch_path(batch_id).parent / "reviews" / batch_id
    paths = sorted(folder.glob("*.json"), key=lambda p: p.stat().st_mtime_ns, reverse=True)
    artifacts = []
    for path in paths[offset:offset + limit]:
        report = json.loads(path.read_text(encoding="utf-8"))
        artifacts.append({key: report[key] for key in
                          ("report_id", "batch_id", "period", "status", "parent_session_id")}
                         | {"created_at": path.stat().st_mtime})
    return {"batch_id": batch_id, "total": len(paths), "offset": offset, "artifacts": artifacts}
