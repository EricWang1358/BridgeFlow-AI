"""Four department sheets into the 跨部门业务整合总表, exactly as the business dictionary says.

The business side supplied the standard templates, the dictionary (字典 v2) and the master
template on 2026-09-13. `data/company_templates/integration.yaml` transcribes the
dictionary: which department column feeds each master field, the join keys, the
formulas it states, and the one cross-department check it names. This module executes
that declaration and nothing else:

- **Rows meet on the declared grain** (项目编号, 客户代码, 报表年月). Names that the
  dictionary says must agree across departments are checked, never merged by similarity.
- **A formula the dictionary states is computed**, and if a department already filled
  that column, the two are compared. A disagreement is reported with both numbers; the
  cell is left empty rather than picking one.
- **What the dictionary does not state is not invented by code.** A formula needing an
  undeclared constant is refused with the constant's name; several rows for one key in one
  department are refused unless that department declares a roll-up. Where the declaration
  fills a gap with an industry convention rather than the business side's own words (the
  VAT rate, the 缺口 formula, the diagnosis thresholds, daily roll-up), it says so under
  `assumptions`, and every cell that depends on one carries that note.
- **Every cell says where it came from**: department, file, sheet, row and header, or the
  formula and the fields it used.

No business name is written in this code; all of them come from the declaration.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal

import openpyxl
import yaml
from fastapi import HTTPException
from pydantic import BaseModel, ConfigDict, Field, model_validator

from bridgeflow import business
from bridgeflow.config import REPO_ROOT, settings


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class SourceColumn(_Strict):
    column: str | None = None
    #: Which occurrence of a repeated header (1-based), for templates that repeat one label.
    occurrence: int = Field(1, ge=1)
    #: A period assembled from a year column and a month column, in that order.
    period_from: list[str] | None = None

    @model_validator(mode="after")
    def _one_kind(self) -> SourceColumn:
        if (self.column is None) == (self.period_from is None):
            raise ValueError("a source names either a column or period_from")
        return self


class FieldDecl(_Strict):
    type: Literal["string", "number"]
    sources: dict[str, SourceColumn]


class Check(_Strict):
    id: str
    title: str
    left: str
    right: str
    tolerance: float = 0


Comparison = Literal["lt", "lte", "gt", "gte"]


class Rule(_Strict):
    label: str
    #: All conditions must hold: [field, comparison, number].
    when: list[tuple[str, Comparison, float]] = Field(min_length=1)


class Classification(_Strict):
    """A text field decided by the first matching rule over numeric fields."""
    rules: list[Rule] = Field(min_length=1)
    otherwise: str


class Rollup(_Strict):
    """How one department's several rows for one key become one. Unlisted fields must be identical."""
    sum: list[str] = Field(default_factory=list)
    concat: list[str] = Field(default_factory=list)
    #: Derived or classified fields: row-level values are dropped and recomputed from the rolled-up inputs.
    recompute: list[str] = Field(default_factory=list)


class ComparisonThreshold(_Strict):
    """When a period-on-period change is worth a reader's attention (E13-UC02)."""
    relative_below: float | None = None
    relative_above: float | None = None


class ComparisonDecl(_Strict):
    """What may be compared across periods, and when a change needs attention.

    `additive` lists the fields that may be summed across entities: quantities and
    amounts. A unit price or a ratio is not additive, so totalling it would produce a
    number with no meaning; those stay per entity.
    """
    additive: list[str] = Field(default_factory=list)
    thresholds: dict[str, ComparisonThreshold] = Field(default_factory=dict)


class CarryOverRule(_Strict):
    """One template column continued from a declared field of an earlier period (E14-UC02)."""
    column: str
    from_field: str
    #: How many periods back, negative. -1 is last month.
    offset: int = Field(default=-1, ge=-12, le=-1)


class DepartmentCarryOver(_Strict):
    rules: list[CarryOverRule] = Field(default_factory=list)


class DepartmentDecl(_Strict):
    label: str
    template: str


class IntegrationSpec(_Strict):
    version: str
    source_dictionary: str
    master_template: str
    departments: dict[str, DepartmentDecl]
    grain: list[str]
    agree: list[str] = Field(default_factory=list)
    constants: dict[str, float | None] = Field(default_factory=dict)
    fields: dict[str, FieldDecl]
    derived: dict[str, dict[str, Any]] = Field(default_factory=dict)
    classifications: dict[str, Classification] = Field(default_factory=dict)
    rollup: dict[str, Rollup] = Field(default_factory=dict)
    checks: list[Check] = Field(default_factory=list)
    comparison: ComparisonDecl = Field(default_factory=ComparisonDecl)
    #: Which template column of a department continues which declared field of an earlier period.
    carry_over: dict[str, DepartmentCarryOver] = Field(default_factory=dict)
    #: Gaps filled by convention, keyed by constant, field, or `rollup.<department>`.
    assumptions: dict[str, str] = Field(default_factory=dict)
    undeclared_rules: dict[str, str] = Field(default_factory=dict)

    @model_validator(mode="after")
    def _references(self) -> IntegrationSpec:
        problems = [f"grain field {g} is not declared" for g in self.grain if g not in self.fields]
        problems += [f"agree field {a} is not declared" for a in self.agree if a not in self.fields]
        for name, decl in self.fields.items():
            problems += [f"{name} names unknown department {d}" for d in decl.sources if d not in self.departments]
        for name, tree in self.derived.items():
            if name not in self.fields:
                problems.append(f"derived {name} is not a declared field")
            for key in _inputs(tree):
                if key not in self.fields:
                    problems.append(f"derived {name} uses undeclared field {key}")
            for key in _declared_constants(tree):
                if key not in self.constants:
                    problems.append(f"derived {name} uses constant {key} not listed in constants")
        for name, table in self.classifications.items():
            if self.fields.get(name) is None or self.fields[name].type != "string":
                problems.append(f"classification {name} is not a declared text field")
            for rule in table.rules:
                problems += [f"classification {name} uses non-numeric field {f}" for f, _, _ in rule.when
                             if self.fields.get(f) is None or self.fields[f].type != "number"]
        for department, policy in self.rollup.items():
            if department not in self.departments:
                problems.append(f"rollup names unknown department {department}")
            for name in policy.sum + policy.concat + policy.recompute:
                if self.fields.get(name) is None or department not in self.fields[name].sources:
                    problems.append(f"rollup.{department} lists {name}, which that department does not supply")
            problems += [f"rollup.{department} sums text field {n}" for n in policy.sum
                         if n in self.fields and self.fields[n].type != "number"]
            problems += [f"rollup.{department} recomputes {n}, which has no formula or rule" for n in policy.recompute
                         if n not in self.derived and n not in self.classifications]
            problems += [f"rollup.{department} cannot roll up grain field {n}" for n in policy.sum + policy.concat + policy.recompute
                         if n in self.grain]
        for name in self.comparison.additive:
            if self.fields.get(name) is None or self.fields[name].type != "number":
                problems.append(f"comparison additive {name} is not a declared number field")
        problems += [f"comparison threshold {n} is not listed as additive" for n in self.comparison.thresholds
                     if n not in self.comparison.additive]
        known = set(self.constants) | set(self.fields) | {f"rollup.{d}" for d in self.rollup}
        problems += [f"assumption {k} names nothing declared" for k in self.assumptions if k not in known]
        for department, policy in self.carry_over.items():
            if department not in self.departments:
                problems.append(f"carry_over names unknown department {department}")
                continue
            for rule in policy.rules:
                decl = self.fields.get(rule.from_field)
                if decl is None or department not in decl.sources:
                    problems.append(f"carry_over.{department} continues {rule.from_field}, "
                                    "which that department does not supply")
        for check in self.checks:
            problems += [f"check {check.id} uses undeclared field {f}" for f in (check.left, check.right) if f not in self.fields]
        if problems:
            raise ValueError("; ".join(problems))
        return self


def _walk(tree: Any):
    if isinstance(tree, dict):
        yield tree
        for child in tree.get("args", []):
            yield from _walk(child)


def _inputs(tree: dict) -> list[str]:
    return [node["key"] for node in _walk(tree) if node.get("op") == "input"]


def _declared_constants(tree: dict) -> list[str]:
    return [node["key"] for node in _walk(tree) if node.get("op") == "declared"]


def spec_path() -> Path:
    path = Path(settings.integration_spec_path)
    return path if path.is_absolute() else REPO_ROOT / path


def load_spec(path: Path | None = None) -> IntegrationSpec:
    target = path or spec_path()
    if not target.is_file():
        raise HTTPException(503, "Integration declaration is not configured")
    try:
        return IntegrationSpec.model_validate(yaml.safe_load(target.read_text(encoding="utf-8")))
    except (OSError, UnicodeError, yaml.YAMLError, ValueError, TypeError) as exc:
        raise HTTPException(503, "Integration declaration is invalid; ask its owner to correct it") from exc


# --- reading a department sheet ---------------------------------------------------------


@dataclass
class Sheet:
    department: str
    filename: str
    sheet: str
    header_row: int  # 1-based, as a person counts
    headers: list[str]
    rows: list[list[Any]]  # data rows only, in order
    #: The sheet row each data row came from, 1-based; blank rows in between are skipped.
    row_numbers: list[int] = field(default_factory=list)


def _clean(label: Any) -> str:
    # pandas renames a repeated header to `label.1`; the template's own spelling may carry
    # stray spaces. Neither is part of what the column means.
    return re.sub(r"\.\d+$", "", str(label or "").strip())


def read_sheet(department: str, filename: str, payload: bytes) -> Sheet:
    """The header row is the first row with more than one filled cell; a title row above it is skipped."""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(415, f"{filename}: use the department's XLSX template")
    workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    if len(workbook.worksheets) != 1:
        raise HTTPException(422, f"{filename}: multiple sheets require explicit selection")
    ws = workbook.worksheets[0]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    header_index = next((i for i, r in enumerate(grid) if sum(c not in (None, "") for c in r) > 1), None)
    if header_index is None:
        raise HTTPException(422, f"{filename}: no header row")
    kept = [(i + 1, r) for i, r in enumerate(grid) if i > header_index and any(c not in (None, "") for c in r)]
    return Sheet(department, filename, ws.title, header_index + 1, [_clean(h) for h in grid[header_index]],
                 [r for _, r in kept], [n for n, _ in kept])


def sheet_from_preview(department: str, source: dict) -> Sheet:
    """Keep the original workbook coordinates selected at import, including skipped titles."""
    header = source.get("header_row", 1)
    return Sheet(department, source.get("filename", ""), source.get("sheet", ""), header,
                 [_clean(h) for h in source["columns"]], source["rows"],
                 source.get("row_numbers", [header + i + 1 for i in range(len(source["rows"]))]))


def _locate(sheet: Sheet, ref: SourceColumn) -> list[int] | None:
    labels = ref.period_from or [ref.column]
    found = []
    for label in labels:
        matches = [i for i, h in enumerate(sheet.headers) if h == _clean(label)]
        occurrence = ref.occurrence if ref.column else 1
        if len(matches) < occurrence:
            return None
        found.append(matches[occurrence - 1])
    return found


def period_field(spec: IntegrationSpec) -> str | None:
    """Which declared field carries the report period, read from the declaration.

    The period is the grain field some department supplies as a period (`period_from`); no
    field name is written here, because the real schema is still being negotiated.
    """
    for name in spec.grain:
        decl = spec.fields.get(name)
        if decl and any(ref.period_from for ref in decl.sources.values()):
            return name
    return None


def periods_in(spec: IntegrationSpec, sheet: Sheet) -> set[str]:
    """The report periods one department's own rows state (E14-UC04 AC-4).

    Unreadable or empty cells are left out rather than guessed: a file whose period cannot be
    read is not thereby a file for this month.
    """
    name = period_field(spec)
    ref = spec.fields[name].sources.get(sheet.department) if name else None
    where = _locate(sheet, ref) if ref else None
    if where is None:
        return set()
    found = set()
    for raw in sheet.rows:
        cells = [raw[i] if i < len(raw) else None for i in where]
        value = _period(*cells) if ref.period_from else _period_text(cells[0])
        if value:
            found.add(value)
    return found


def _period_text(value: Any) -> str | None:
    """A period written in one cell, in any of the forms a template uses (2024-07, 2024年7月)."""
    match = re.search(r"(\d{4})\D{0,3}(\d{1,2})", str(value or ""))
    if not match or not 1 <= int(match.group(2)) <= 12:
        return None
    return f"{match.group(1)}-{int(match.group(2)):02d}"


# --- integrating -------------------------------------------------------------------------


class Issue(BaseModel):
    kind: Literal["missing_department", "missing_column", "needs_rollup", "disagreement", "invalid_number",
                  "invalid_period", "undeclared_constant", "derived_mismatch", "check_failed", "cannot_compute",
                  "missing_key"]
    field: str = ""
    key: list[str] = Field(default_factory=list)
    departments: list[str] = Field(default_factory=list)
    message: str


class MasterRow(BaseModel):
    key: list[str]
    values: dict[str, Any]
    provenance: dict[str, dict[str, Any]]
    complete: bool


class MasterResult(BaseModel):
    version: str
    columns: list[str]
    rows: list[MasterRow]
    issues: list[Issue]
    departments_read: dict[str, dict[str, Any]]
    #: Conventions the declaration relies on in place of a business-confirmed rule.
    assumptions: dict[str, str] = Field(default_factory=dict)


def _period(year: Any, month: Any) -> str | None:
    y = re.search(r"\d{4}", str(year or ""))
    m = re.search(r"\d{1,2}", str(month or "").replace(str(year or ""), "") if month is not None else "")
    if not y or not m or not 1 <= int(m.group()) <= 12:
        return None
    return f"{y.group()}-{int(m.group()):02d}"


def _number(value: Any) -> Decimal | None:
    if isinstance(value, bool):
        return None
    text = str(value).strip().replace(",", "")
    percent = text.endswith("%")
    try:
        number = Decimal(text.rstrip("%"))
    except InvalidOperation:
        return None
    if not number.is_finite():
        return None
    return number / 100 if percent else number


def _plain(number: Decimal) -> int | float:
    return int(number) if number == number.to_integral_value() else float(number)


@dataclass
class _Contribution:
    sheet: Sheet
    row_number: int  # 1-based sheet row
    values: dict[str, Any] = field(default_factory=dict)
    #: Set when several rows were rolled up: every sheet row used, and each field's policy.
    rows: list[int] = field(default_factory=list)
    policies: dict[str, str] = field(default_factory=dict)


def integrate(spec: IntegrationSpec, sheets: list[Sheet]) -> MasterResult:
    issues: list[Issue] = []
    columns = list(spec.fields)
    by_department: dict[str, dict[tuple[str, ...], list[_Contribution]]] = {}
    read: dict[str, dict[str, Any]] = {}

    for sheet in sheets:
        if sheet.department not in spec.departments:
            raise HTTPException(422, f"Unknown department {sheet.department}")
        located: dict[str, list[int]] = {}
        missing = []
        for name, decl in spec.fields.items():
            ref = decl.sources.get(sheet.department)
            if ref is None:
                continue
            where = _locate(sheet, ref)
            if where is None:
                missing.append(f"{name}←{'+'.join(ref.period_from or [ref.column or ''])}")
            else:
                located[name] = where
        if missing:
            issues.append(Issue(kind="missing_column", departments=[sheet.department],
                                message=f"{spec.departments[sheet.department].label} template lacks: {', '.join(missing)}"))
        read[sheet.department] = {"filename": sheet.filename, "sheet": sheet.sheet, "header_row": sheet.header_row,
                                  "data_rows": len(sheet.rows), "missing_columns": len(missing)}
        groups: dict[tuple[str, ...], list[_Contribution]] = {}
        for offset, raw in enumerate(sheet.rows):
            number = sheet.row_numbers[offset] if offset < len(sheet.row_numbers) else sheet.header_row + 1 + offset
            contribution = _Contribution(sheet, number)
            for name, where in located.items():
                ref = spec.fields[name].sources[sheet.department]
                cells = [raw[i] if i < len(raw) else None for i in where]
                if ref.period_from:
                    value = _period(*cells)
                    if value is None and any(c not in (None, "") for c in cells):
                        issues.append(Issue(kind="invalid_period", field=name, departments=[sheet.department],
                                            message=f"{sheet.filename} row {contribution.row_number}: year and month do not form a period"))
                else:
                    value = cells[0]
                if value not in (None, ""):
                    contribution.values[name] = value
            key = tuple(str(contribution.values.get(g, "")).strip() for g in spec.grain
                        if spec.fields[g].sources.get(sheet.department))
            if not all(key) or len(key) != sum(1 for g in spec.grain if spec.fields[g].sources.get(sheet.department)):
                issues.append(Issue(kind="missing_key", departments=[sheet.department],
                                    message=f"{sheet.filename} row {contribution.row_number}: a grain field is empty; the row cannot be placed"))
                continue
            groups.setdefault(key, []).append(contribution)
        by_department[sheet.department] = groups

    keys = sorted({k for groups in by_department.values() for k in groups})
    rows: list[MasterRow] = []
    for key in keys:
        values: dict[str, Any] = dict.fromkeys(columns)
        provenance: dict[str, dict[str, Any]] = {}
        complete = True
        present = []
        variants: dict[str, dict[str, list[str]]] = {}
        for department, groups in by_department.items():
            group = groups.get(key)
            if not group:
                continue
            if len(group) > 1:
                merged = _roll_up(spec, department, group)
                if isinstance(merged, Issue):
                    complete = False
                    merged.key = list(key)
                    issues.append(merged)
                    continue
                group = [merged]
            present.append(department)
            for name, value in group[0].values.items():
                decl = spec.fields[name]
                where = {"department": department, "file": group[0].sheet.filename, "sheet": group[0].sheet.sheet,
                         "row": group[0].row_number, "column": "+".join(decl.sources[department].period_from or [decl.sources[department].column or ""])}
                if group[0].rows:
                    where |= {"rows": group[0].rows, "rollup": group[0].policies.get(name, "identical")}
                if decl.type == "number":
                    number = _number(value)
                    if number is None:
                        complete = False
                        issues.append(Issue(kind="invalid_number", field=name, key=list(key), departments=[department],
                                            message=f"{name}: {where['file']} row {where['row']} is not a number"))
                        continue
                    value = _plain(number)
                variants.setdefault(name, {}).setdefault(str(value).strip(), []).append(department)
                if values[name] is None:
                    values[name], provenance[name] = value, where
        for name, written in variants.items():
            if len(written) > 1:
                # One question per field, whatever order the departments arrived in, and
                # no department's spelling is picked for the cell.
                complete = False
                values[name] = None
                provenance[name] = {"conflict": written}
                issues.append(Issue(kind="disagreement", field=name, key=list(key),
                                    departments=sorted({d for ds in written.values() for d in ds}),
                                    message=f"{name} differs between departments ("
                                            + "; ".join(f"{'/'.join(ds)}: {v}" for v, ds in written.items())
                                            + "); confirm which is right"))
        absent = [d for d in spec.departments if d not in present]
        if absent:
            complete = False
            issues.append(Issue(kind="missing_department", key=list(key), departments=absent,
                                message="No row for this key from: " + ", ".join(spec.departments[d].label for d in absent)))

        for name, tree in spec.derived.items():
            outcome = _derive(spec, name, tree, values)
            supplied = values.get(name)
            if isinstance(outcome, Issue):
                complete = False
                # The formula cannot run. A value the department wrote is kept as written and
                # marked unverified; nothing is filled in where nobody wrote anything.
                outcome.key = list(key)
                issues.append(outcome)
                if supplied is not None:
                    provenance[name] = {**provenance[name], "verified_by_formula": False, "unverified_because": outcome.message}
                continue
            computed, formula_inputs = outcome
            if supplied is None:
                values[name] = computed
                provenance[name] = {"formula": name, "inputs": formula_inputs}
            elif abs(Decimal(str(supplied)) - Decimal(str(computed))) > _written_tolerance(supplied, computed):
                complete = False
                issues.append(Issue(kind="derived_mismatch", field=name, key=list(key), departments=[provenance[name]["department"]],
                                    message=f"{name}: department wrote {supplied}, the dictionary formula gives {computed}"))
                values[name] = None
                provenance[name] = {"conflict": {"supplied": supplied, "computed": computed}}
            else:
                provenance[name] = {**provenance[name], "verified_by_formula": True}

        for name, table in spec.classifications.items():
            outcome = _classify(name, table, values)
            supplied = values.get(name)
            if isinstance(outcome, Issue):
                complete = False
                outcome.key = list(key)
                issues.append(outcome)
                if supplied is not None:
                    provenance[name] = {**provenance[name], "verified_by_rule": False, "unverified_because": outcome.message}
                continue
            label, rule_inputs = outcome
            if supplied is None:
                values[name] = label
                provenance[name] = {"rule": name, "inputs": rule_inputs}
            elif str(supplied).strip() != label:
                complete = False
                issues.append(Issue(kind="derived_mismatch", field=name, key=list(key), departments=[provenance[name]["department"]],
                                    message=f"{name}: department wrote {supplied}, the declared rule gives {label}"))
                values[name] = None
                provenance[name] = {"conflict": {"supplied": supplied, "computed": label}}
            else:
                provenance[name] = {**provenance[name], "verified_by_rule": True}

        for name, cell in provenance.items():
            notes = _assumptions_behind(spec, name, cell)
            if notes:
                cell["assumptions"] = notes

        for check in spec.checks:
            left, right = values.get(check.left), values.get(check.right)
            if left is None or right is None:
                continue
            if abs(Decimal(str(left)) - Decimal(str(right))) > Decimal(str(check.tolerance)):
                complete = False
                issues.append(Issue(kind="check_failed", field=check.left, key=list(key),
                                    message=f"{check.title}: {check.left}={left}, {check.right}={right}"))
        rows.append(MasterRow(key=list(key), values=values, provenance=provenance, complete=complete))

    return MasterResult(version=spec.version, columns=columns, rows=rows, issues=issues, departments_read=read,
                        assumptions=spec.assumptions)


def _roll_up(spec: IntegrationSpec, department: str, group: list[_Contribution]) -> _Contribution | Issue:
    label = spec.departments[department].label
    policy = spec.rollup.get(department)
    if policy is None:
        return Issue(kind="needs_rollup", departments=[department],
                     message=f"{label} has {len(group)} rows for this key and no roll-up rule is declared")
    merged = _Contribution(group[0].sheet, group[0].row_number, rows=[c.row_number for c in group])
    conflicting = []
    for name in dict.fromkeys(n for c in group for n in c.values):
        present = [c.values[name] for c in group if name in c.values]
        if name in policy.recompute:
            merged.policies[name] = "recomputed"
        elif name in policy.sum:
            if len(present) != len(group):
                return Issue(kind="invalid_number", field=name, departments=[department],
                             message=f"{name}: a row has no value; a partial sum is not a total")
            numbers = [_number(v) for v in present]
            invalid = next((v for v, n in zip(present, numbers, strict=True) if n is None), None)
            merged.values[name] = invalid if invalid is not None else _plain(sum(numbers, Decimal(0)))
            merged.policies[name] = "sum"
        elif name in policy.concat:
            merged.values[name] = "；".join(dict.fromkeys(str(v).strip() for v in present))
            merged.policies[name] = "concat"
        else:
            numeric = spec.fields[name].type == "number"
            distinct = {(_number(v) if numeric and _number(v) is not None else str(v).strip()) for v in present}
            if len(distinct) > 1:
                conflicting.append(name)
            merged.values[name] = present[0]
    if conflicting:
        return Issue(kind="needs_rollup", field=conflicting[0], departments=[department],
                     message=f"{label} has {len(group)} rows for this key that differ in {', '.join(conflicting)}, which the roll-up does not sum, join or recompute")
    return merged


_COMPARE = {"lt": lambda a, b: a < b, "lte": lambda a, b: a <= b, "gt": lambda a, b: a > b, "gte": lambda a, b: a >= b}


def _classify(name: str, table: Classification, values: dict[str, Any]) -> tuple[str, list[str]] | Issue:
    needed = list(dict.fromkeys(f for rule in table.rules for f, _, _ in rule.when))
    missing = [f for f in needed if values.get(f) is None]
    if missing:
        return Issue(kind="cannot_compute", field=name, message=f"{name} needs {', '.join(missing)}")
    for rule in table.rules:
        if all(_COMPARE[op](Decimal(str(values[f])), Decimal(str(bound))) for f, op, bound in rule.when):
            return rule.label, needed
    return table.otherwise, needed


def _assumptions_behind(spec: IntegrationSpec, name: str, cell: dict[str, Any]) -> list[str]:
    keys = [name] if "formula" in cell or "rule" in cell or cell.get("verified_by_formula") is not None \
        or cell.get("verified_by_rule") is not None else []
    if name in spec.derived and keys:
        keys += _declared_constants(spec.derived[name])
    if "rollup" in cell:
        keys.append(f"rollup.{cell['department']}")
    return [f"{k}: {spec.assumptions[k]}" for k in dict.fromkeys(keys) if k in spec.assumptions]


def _written_tolerance(supplied: Any, computed: Any) -> Decimal:
    """How far a department's figure may sit from the formula: half a unit of the last digit it wrote.

    People round — a sign-off rate written as 99.12% is right if the formula gives 0.991234.
    A figure written to full precision is held to a relative 1e-6.
    """
    exponent = Decimal(str(supplied)).normalize().as_tuple().exponent
    written = Decimal(5).scaleb(exponent - 1) if isinstance(exponent, int) and -8 <= exponent <= 0 else Decimal(0)
    return max(written, Decimal("1e-6") * max(Decimal(1), abs(Decimal(str(computed)))))


def _derive(spec: IntegrationSpec, name: str, tree: dict, values: dict[str, Any]) -> tuple[int | float, list[str]] | Issue:
    for constant in _declared_constants(tree):
        if spec.constants.get(constant) is None:
            return Issue(kind="undeclared_constant", field=name,
                         message=f"{name} needs the undeclared constant 「{constant}」")
    resolved = _substitute(tree, spec.constants)
    needed = _inputs(tree)
    missing = [k for k in needed if values.get(k) is None]
    if missing:
        return Issue(kind="cannot_compute", field=name, message=f"{name} needs {', '.join(missing)}")
    inputs = {k: business.Value(Decimal(str(values[k])), [], 1) for k in needed}
    try:
        result = business.expression(resolved, None, inputs=inputs).number
    except HTTPException as exc:
        return Issue(kind="cannot_compute", field=name, message=f"{name}: {exc.detail}")
    return _plain(result), needed


def _substitute(tree: Any, constants: dict[str, float | None]) -> Any:
    if not isinstance(tree, dict):
        return tree
    if tree.get("op") == "declared":
        return {"op": "constant", "value": constants[tree["key"]]}
    return {**tree, **({"args": [_substitute(a, constants) for a in tree["args"]]} if "args" in tree else {})}


def to_xlsx(result: MasterResult) -> bytes:
    """The master template's columns, one row per key; issues on a second sheet."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "总表"
    sheet.append(result.columns)
    for row in result.rows:
        sheet.append([row.values.get(c) for c in result.columns])
    notes = book.create_sheet("待确认")
    notes.append(["类型", "字段", "键", "部门", "说明"])
    for issue in result.issues:
        notes.append([issue.kind, issue.field, " / ".join(issue.key), ", ".join(issue.departments), issue.message])
    assumed = book.create_sheet("口径假设")
    assumed.append(["规则", "说明"])
    for name, text in result.assumptions.items():
        assumed.append([name, text])
    # This export contains values, never executable formulas, including labels and notes.
    for worksheet in book:
        for row in worksheet:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
