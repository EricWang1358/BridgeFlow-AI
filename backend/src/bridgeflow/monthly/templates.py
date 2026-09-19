"""The department template for a month, prefilled from last month where declared (E14-UC02).

Every month the same person fills the same sheet, and part of it is last month's numbers
copied by hand. Copying by hand is where a month's figures quietly drift from the month they
claim to continue, so the carry-over is declared (`carry_over` in the integration
declaration) and done here: which template column continues which declared field, and from
how many periods back.

Three things this deliberately does **not** do:

- It does not invent a column. The workbook's header is the approved template's header,
  copied cell for cell, so what comes back can be imported by the same declaration that
  produced it.
- It does not prefill from a value nobody could see: the prior period's figures come from
  the master of that period's batch, through the same visibility the reader has.
- It does not hide that a number was carried over. Each prefilled cell is marked with the
  period it came from and asks to be checked, and the instructions sheet says the same in
  words — a prefilled cell is a draft, not a measurement.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

#: A light fill marks a prefilled cell; the comment and the instructions sheet carry the
#: meaning, because colour alone is not a statement (00-foundations §5.5).
PREFILL_FILL = PatternFill("solid", fgColor="FFF2CC")


class CarryOverRule(BaseModel):
    """One template column continued from a declared field of an earlier period."""
    column: str
    from_field: str
    #: How many periods back, as a negative number. -1 is last month.
    offset: int = Field(default=-1, le=-1, ge=-12)


class DepartmentCarryOver(BaseModel):
    rules: list[CarryOverRule] = Field(default_factory=list)


def declared(spec: dict, department: str) -> list[CarryOverRule]:
    raw = ((spec or {}).get("carry_over") or {}).get(department) or {}
    return DepartmentCarryOver.model_validate(raw).rules


def _key_columns(spec: dict, department: str) -> dict[str, str]:
    """Which template columns carry this department's part of the grain, per the declaration."""
    columns = {}
    for name in (spec.get("grain") or []):
        source = ((spec.get("fields") or {}).get(name) or {}).get("sources", {}).get(department)
        if source and source.get("column"):
            columns[name] = source["column"]
    return columns


def _field_column(spec: dict, department: str, field: str) -> str:
    source = ((spec.get("fields") or {}).get(field) or {}).get("sources", {}).get(department) or {}
    return source.get("column") or ""


def build(*, spec: dict, department: str, period: str, template: bytes,
          prior: Any = None, prior_period: str = "", axis: list[int] | None = None,
          reason: str = "") -> bytes:
    """The workbook to hand a department for `period`.

    `prior` is the master result of the period the carry-over reads, or None. `reason` says
    why there is none, and is printed rather than left to be guessed at.
    """
    book = openpyxl.load_workbook(io.BytesIO(template))
    sheet = book.worksheets[0]
    header = [cell.value for cell in sheet[1]]
    rules = [rule for rule in declared(spec, department) if rule.column in header]
    keys = _key_columns(spec, department)
    carried = 0
    if prior is not None and rules:
        positions = {str(name): index + 1 for index, name in enumerate(header) if name}
        seen: set[tuple[str, ...]] = set()
        row_number = 2
        for row in prior.rows:
            entity = tuple(row.key[i] for i in (axis or range(len(row.key))) if i < len(row.key))
            if entity in seen:
                continue
            seen.add(entity)
            wrote = False
            for rule in rules:
                value = row.values.get(rule.from_field)
                if value is None:
                    continue
                cell = sheet.cell(row=row_number, column=positions[rule.column], value=value)
                cell.fill = PREFILL_FILL
                cell.comment = Comment(f"预填自 {prior_period}（{rule.from_field}），请核对 / "
                                       f"Carried over from {prior_period}; please check", "BridgeFlow")
                wrote = True
            if not wrote:
                continue
            for name, column in keys.items():
                if column in positions and name in (spec.get("fields") or {}):
                    index = (spec.get("grain") or []).index(name)
                    if index < len(row.key):
                        sheet.cell(row=row_number, column=positions[column], value=row.key[index])
            carried += 1
            row_number += 1
    _instructions(book, spec=spec, department=department, period=period, rules=rules,
                  prior_period=prior_period if prior is not None else "", carried=carried, reason=reason)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _instructions(book, *, spec: dict, department: str, period: str, rules: list[CarryOverRule],
                  prior_period: str, carried: int, reason: str) -> None:
    """The sheet that says what each column means, what was prefilled and what was not."""
    sheet = book.create_sheet("填写说明")
    bold = Font(bold=True)
    lines: list[list[Any]] = [
        ["月度模板填写说明 / How to fill this template"],
        ["期间 / Period", period],
        ["部门 / Department", (spec.get("departments") or {}).get(department, {}).get("label", department)],
        ["声明版本 / Declaration", spec.get("version", "")],
        [],
        ["预填 / Carried over"],
    ]
    if rules and prior_period:
        lines.append([f"已按声明从 {prior_period} 预填 {carried} 行；预填单元格带批注，请逐行核对后再提交。"])
        lines.append([f"Prefilled {carried} row(s) from {prior_period}. Prefilled cells carry a note; check each one."])
        for rule in rules:
            lines.append([rule.column, f"← {rule.from_field}（{prior_period}）"])
    else:
        lines.append([reason or "没有可用的上期批次，本次不预填任何沿用字段。"])
        lines.append([reason or "No usable prior batch; nothing was prefilled."])
    lines += [[], ["字段说明 / Fields"], ["列 / Column", "类型 / Type", "对应声明字段 / Declared field", "必填 / Required"]]
    grain = set(spec.get("grain") or [])
    for name, decl in (spec.get("fields") or {}).items():
        source = (decl.get("sources") or {}).get(department)
        if not source or not source.get("column"):
            continue
        lines.append([source["column"], decl.get("type", ""), name, "是 / yes" if name in grain else ""])
    for index, line in enumerate(lines, start=1):
        for column, value in enumerate(line, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            if column == 1 and len(line) == 1:
                cell.font = bold
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 28
