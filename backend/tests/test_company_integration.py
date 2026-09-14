"""The business side's templates and dictionary (2026-09-13) drive the master table.

The golden case is theirs: the worked example row in `master-v2.xlsx`, split back into
the four department templates by `scripts/make_template_example.py`. Every other test
changes one thing in those sheets and checks the system refuses to guess.
"""
import ast
import io
from pathlib import Path

import openpyxl

from bridgeflow import integration
from bridgeflow.config import REPO_ROOT

BASE = REPO_ROOT / "data" / "company_templates"
SPEC = integration.load_spec(BASE / "integration.yaml")


def workbook_bytes(department: str, edit=None, extra_rows=0) -> bytes:
    book = openpyxl.load_workbook(BASE / "example" / f"{department}.xlsx")
    sheet = book.active
    if edit:
        edit(sheet)
    for _ in range(extra_rows):
        sheet.append([c.value for c in sheet[2]])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def column(sheet, label, occurrence=1):
    matches = [c.column_letter for c in sheet[1] if str(c.value or "").strip() == label]
    return matches[occurrence - 1]


def sheets(**changes):
    result = []
    for department in SPEC.departments:
        if changes.get(department) == "absent":
            continue
        edit, extra = changes.get(department, (None, 0))
        result.append(integration.read_sheet(department, f"{department}.xlsx", workbook_bytes(department, edit, extra)))
    return result


def expected_row():
    master = openpyxl.load_workbook(BASE / SPEC.master_template).active
    return dict(zip([c.value for c in master[1]], [c.value for c in master[2]], strict=True))


def kinds(result):
    return sorted({i.kind for i in result.issues})


# --- the declaration is a faithful transcription ---------------------------------------


def test_declared_fields_are_the_dictionary_and_the_master_template_in_order():
    dictionary = openpyxl.load_workbook(BASE / SPEC.source_dictionary).active
    names = [r[0] for r in dictionary.iter_rows(min_row=2, values_only=True) if r[0]]
    master = [c.value for c in openpyxl.load_workbook(BASE / SPEC.master_template).active[1]]
    assert list(SPEC.fields) == names == master


def test_every_declared_source_column_exists_in_its_department_template():
    for department, decl in SPEC.departments.items():
        headers = [str(c.value or "").strip() for c in openpyxl.load_workbook(BASE / decl.template).active[1]]
        for name, field in SPEC.fields.items():
            ref = field.sources.get(department)
            if ref is None:
                continue
            for label in ref.period_from or [ref.column]:
                count = headers.count(label)
                assert count >= (ref.occurrence if ref.column else 1), f"{department}: {name} → {label} not in template"


# --- golden case --------------------------------------------------------------------------


def test_the_business_example_is_reproduced_column_for_column():
    result = integration.integrate(SPEC, sheets())
    [row] = result.rows
    assert row.values == expected_row()
    assert row.complete
    assert result.issues == []
    assert row.provenance["生产_实际签收率"]["verified_by_formula"] is True
    assert row.provenance["物资_单方不含税毛利"]["verified_by_formula"] is True
    assert row.provenance["市场_缺口"]["verified_by_formula"] is True
    assert row.provenance["生产_客户合作状态诊断"]["verified_by_rule"] is True
    assert row.provenance["客户名称"]["column"] in {"客户名称", "客户单位"}


def test_without_a_declared_vat_rate_the_margin_is_kept_but_unverified():
    undeclared = SPEC.model_copy(update={"constants": {"增值税税率": None}})
    result = integration.integrate(undeclared, sheets())
    assert kinds(result) == ["undeclared_constant"]
    assert result.rows[0].provenance["物资_单方不含税毛利"]["verified_by_formula"] is False
    assert result.rows[0].values["物资_单方不含税毛利"] == expected_row()["物资_单方不含税毛利"]


def test_every_cell_resting_on_a_convention_says_which_one():
    [row] = integration.integrate(SPEC, sheets()).rows
    assert any(n.startswith("增值税税率") for n in row.provenance["物资_单方不含税毛利"]["assumptions"])
    assert any(n.startswith("市场_缺口") for n in row.provenance["市场_缺口"]["assumptions"])
    assert "assumptions" not in row.provenance["生产_实际签收率"]  # stated by the dictionary itself
    assert "assumptions" not in row.provenance["财务_期初金额"]


def test_the_diagnosis_follows_the_declared_rule_table():
    def month(signed, previous):
        def edit(sheet):
            sheet[f"{column(sheet, '实际签收率')}2"] = None
            sheet[f"{column(sheet, '产量环比增长率')}2"] = None
            sheet[f"{column(sheet, '客户合作状态诊断')}2"] = None
            sheet[f"{column(sheet, '实际量')}2"] = signed
            sheet[f"{column(sheet, '上月实际量')}2"] = previous
        [row] = integration.integrate(SPEC, sheets(production=(edit, 0))).rows
        return row.values["生产_客户合作状态诊断"]
    # 出厂量 is 4980 in the example.
    assert month(4500, 4500) == "签收异常"      # 签收率 90.4%
    assert month(4970, 4800) == "平稳合作"      # +3.5%
    assert month(4970, 5400) == "需求下滑"      # −8.0%
    assert month(4970, 6300) == "合作萎缩"      # −21.1%
    assert month(4970, 4600) == "稳定增长"      # +8.0%


def test_a_diagnosis_that_contradicts_the_rule_is_withheld():
    def wrong(sheet):
        sheet[f"{column(sheet, '客户合作状态诊断')}2"] = "合作萎缩"
    result = integration.integrate(SPEC, sheets(production=(wrong, 0)))
    [issue] = [i for i in result.issues if i.kind == "derived_mismatch"]
    assert issue.field == "生产_客户合作状态诊断" and result.rows[0].values[issue.field] is None


# --- refusing to guess -------------------------------------------------------------------


def test_a_department_figure_that_contradicts_its_formula_is_withheld_with_both_numbers():
    def wrong_rate(sheet):
        sheet[f"{column(sheet, '实际签收率')}2"] = 0.5
    result = integration.integrate(SPEC, sheets(production=(wrong_rate, 0)))
    [mismatch] = [i for i in result.issues if i.kind == "derived_mismatch"]
    assert mismatch.field == "生产_实际签收率" and "0.5" in mismatch.message
    assert result.rows[0].values["生产_实际签收率"] is None and not result.rows[0].complete


def test_the_declared_cross_department_check_catches_different_volumes():
    def other_volume(sheet):
        sheet[f"{column(sheet, '当月生产量')}2"] = 4000
    result = integration.integrate(SPEC, sheets(procurement=(other_volume, 0)))
    assert "check_failed" in kinds(result)


def test_names_that_must_agree_are_checked_not_merged():
    def renamed(sheet):
        sheet[f"{column(sheet, '项目名称')}2"] = "项目B"
    result = integration.integrate(SPEC, sheets(marketing=(renamed, 0)))
    [issue] = [i for i in result.issues if i.kind == "disagreement"]
    assert issue.field == "项目名称"


def test_a_missing_department_leaves_a_partial_row_that_says_so():
    result = integration.integrate(SPEC, sheets(finance="absent"))
    assert result.rows[0].complete is False
    [issue] = [i for i in result.issues if i.kind == "missing_department"]
    assert issue.departments == ["finance"]
    assert result.rows[0].values["财务_期初金额"] is None


def daily(splits):
    """The example's production row split into days; each day carries its own ratios."""
    def edit(sheet):
        base = [c.value for c in sheet[2]]
        headers = [str(c.value or "").strip() for c in sheet[1]]
        sheet.delete_rows(2)
        for share, station in splits:
            row = list(base)
            for label in ("生产量", "出厂量", "实际量"):
                row[headers.index(label)] = base[headers.index(label)] * share
            row[headers.index("实际签收率")] = 0.5  # a day's own ratio, which must not be averaged or kept
            row[headers.index("厂站")] = station
            sheet.append(row)
    return edit


def test_daily_production_rows_roll_up_to_the_business_example():
    result = integration.integrate(SPEC, sheets(production=(daily([(0.4, "A拌站"), (0.6, "A拌站")]), 0)))
    [row] = result.rows
    assert result.issues == [] and row.complete
    assert row.values == expected_row()
    assert row.provenance["生产_生产量"]["rows"] == [2, 3] and row.provenance["生产_生产量"]["rollup"] == "sum"
    assert row.provenance["生产_实际签收率"]["formula"] == "生产_实际签收率"
    assert any(n.startswith("rollup.production") for n in row.provenance["生产_生产量"]["assumptions"])


def test_roll_up_joins_distinct_stations():
    [row] = integration.integrate(SPEC, sheets(production=(daily([(0.5, "A拌站"), (0.5, "B拌站")]), 0))).rows
    assert row.values["生产_厂站"] == "A拌站；B拌站"


def test_rows_that_differ_in_a_field_the_roll_up_does_not_cover_are_refused():
    def edit(sheet):
        daily([(0.5, "A拌站"), (0.5, "A拌站")])(sheet)
        sheet[f"{column(sheet, '上月实际量')}3"] = 1
    result = integration.integrate(SPEC, sheets(production=(edit, 0)))
    [issue] = [i for i in result.issues if i.kind == "needs_rollup"]
    assert issue.field == "生产_上月实际量"
    assert result.rows[0].values["生产_生产量"] is None


def test_a_department_without_a_declared_roll_up_still_refuses_several_rows():
    result = integration.integrate(SPEC, sheets(finance=(None, 1)))
    assert "needs_rollup" in kinds(result)
    assert result.rows[0].values["财务_期初金额"] is None


def test_the_v1_templates_lack_the_join_keys_the_dictionary_added_in_v2():
    """字典 v2 upgraded 项目编号/客户代码 to keys for all four departments; v1 sheets cannot be placed."""
    v1 = [integration.read_sheet(d, f"{d}-v1.xlsx", (BASE / "source" / f"{d}-v1.xlsx").read_bytes()) for d in SPEC.departments]
    result = integration.integrate(SPEC, v1)
    assert result.rows == []
    assert {i.departments[0] for i in result.issues if i.kind == "missing_column"} == set(SPEC.departments)
    procurement = result.departments_read["procurement"]
    assert procurement["header_row"] == 2  # the v1 sheet has a merged title row above its header


def test_the_master_workbook_keeps_template_columns_and_lists_open_items():
    result = integration.integrate(SPEC, sheets())
    book = openpyxl.load_workbook(io.BytesIO(integration.to_xlsx(result)))
    assert [c.value for c in book["总表"][1]] == list(SPEC.fields)
    assert book["待确认"].max_row == 1 + len(result.issues)
    assert book["口径假设"].max_row == 1 + len(SPEC.assumptions)


def test_the_example_generator_is_reproducible(tmp_path):
    import subprocess
    import sys
    before = {p.name: p.read_bytes() for p in (BASE / "example").glob("*.xlsx")}
    subprocess.run([sys.executable, str(REPO_ROOT / "scripts" / "make_template_example.py"), str(tmp_path)], check=True, capture_output=True)
    after = {name: openpyxl.load_workbook(tmp_path / name).active for name in before}
    for name, sheet in after.items():
        original = openpyxl.load_workbook(io.BytesIO(before[name])).active
        assert [[c.value for c in r] for r in sheet.iter_rows()] == [[c.value for c in r] for r in original.iter_rows()]


def test_no_business_name_is_written_into_the_integration_code():
    tree = ast.parse(Path(integration.__file__).read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        holder = isinstance(node, ast.Module | ast.ClassDef | ast.FunctionDef | ast.AsyncFunctionDef)
        if holder and ast.get_docstring(node) is not None:
            node.body = node.body[1:] or [ast.Pass()]
    code = ast.unparse(tree)
    for name in [*SPEC.fields, *SPEC.constants]:
        assert name not in code, name


def test_an_imported_batch_of_the_templates_produces_the_master_table_over_http(monkeypatch):
    import base64

    from fastapi.testclient import TestClient

    from bridgeflow.api.main import app
    from bridgeflow.config import settings
    monkeypatch.setattr(settings, "field_dictionary_path", str(REPO_ROOT / "data/business_demo/dictionary.yaml"))
    with TestClient(app) as client:
        files = [("files", (f"{d}.xlsx", (BASE / "example" / f"{d}.xlsx").read_bytes(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for d in SPEC.departments]
        batch = client.post("/batches", data={"period": "2024-05", "departments": list(SPEC.departments)}, files=files)
        assert batch.status_code == 200, batch.text
        batch_id = batch.json()["batch_id"]
        view = client.get(f"/integration/batches/{batch_id}").json()
        assert view["rows"][0]["values"] == expected_row()
        summary = client.post("/tools/integration-summary", json={"batch_id": batch_id}).json()
        assert (summary["rows"], summary["complete_rows"], summary["issues_by_kind"]) == (1, 1, {})
        assert "项目A" not in str(summary) and "1485000" not in str(summary)
        workbook = client.get(f"/integration/batches/{batch_id}/xlsx").json()
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(workbook["base64"])))
        assert [c.value for c in book["总表"][2]] == list(expected_row().values())


def test_a_rounded_department_figure_is_checked_at_the_precision_it_was_written():
    def rounded(value):
        def edit(sheet):
            sheet[f"{column(sheet, '实际签收率')}2"] = value
        return integration.integrate(SPEC, sheets(production=(edit, 0)))
    assert rounded(0.9940).rows[0].provenance["生产_实际签收率"]["verified_by_formula"] is True   # 0.993976 → 99.40%
    assert rounded("99.4%").rows[0].complete
    assert "derived_mismatch" in kinds(rounded(0.9930))                                           # off by a full unit
