"""The English sample set (data/demo_en/) is the Chinese one in other words, nothing else.

The public demo loads the English translation (docs/37). A translation that changed a number,
merged two columns or lost a planted problem would demo a product that behaves differently
from the one the Chinese tests prove, so every case is imported from both sets and compared
outcome by outcome, and the set is checked against its generator.
"""
import hashlib
import json
import re
import subprocess
import sys

import pytest
import yaml
from fastapi.testclient import TestClient

from bridgeflow.config import REPO_ROOT, settings

SETS = {"zh": "data/mock_business/sample-set.yaml", "en": "data/demo_en/sample-set.yaml"}
CJK = re.compile(r"[　-〿一-鿿＀-￯]")


def use(monkeypatch, name: str) -> None:
    declared = yaml.safe_load((REPO_ROOT / SETS[name]).read_text(encoding="utf-8"))
    for key, setting in (("dictionary", "field_dictionary_path"), ("integration_spec", "integration_spec_path"),
                         ("demo_cases", "demo_cases_path"), ("workflow_catalogue", "workflow_catalogue_path"),
                         ("workflow_samples", "workflow_samples_path"), ("discovery_sample", "discovery_sample_path"),
                         ("discovery_scoring_policy", "discovery_scoring_policy_path"),
                         ("discovery_decision_policy", "discovery_decision_policy_path")):
        monkeypatch.setattr(settings, setting, str(REPO_ROOT / declared[key]))


def outcome(client: TestClient, case: str) -> dict:
    """Everything a case is meant to show, with names left out so the two sets compare."""
    batch = client.post(f"/batches/demo?case={case}")
    assert batch.status_code == 200, batch.text
    body = batch.json()
    master = client.get(f"/integration/batches/{body['batch_id']}").json()
    context = client.post("/tools/review-context", json={"batch_id": body["batch_id"]})
    checks = sorted((c["check_id"], c["expected_status"], c.get("value"))
                    for role in context.json()["roles"] for c in role["checks"]) if context.status_code == 200 else []
    return {
        "status": body["status"], "period": body["period"], "master_rows": body["master_rows"],
        "demo_case": body["demo_case"], "column_questions": body["column_questions"],
        "blockers": len(body["review_blockers"]),
        "departments": sorted((d["department"], d["rows"], d["quarantined"]) for d in body["departments"]),
        "issues": sorted(issue["kind"] for issue in master["issues"]),
        "complete_rows": sum(1 for row in master["rows"] if row.get("complete")),
        "review": context.status_code, "checks": checks,
    }


@pytest.fixture
def both(monkeypatch):
    def run(case: str) -> tuple[dict, dict]:
        results = {}
        for name in SETS:
            use(monkeypatch, name)
            with TestClient(__import__("bridgeflow.api.main", fromlist=["app"]).app) as client:
                results[name] = outcome(client, case)
        return results["zh"], results["en"]
    return run


@pytest.mark.parametrize("case", ["tour", "core", "other", "clean"])
def test_every_case_comes_out_the_same_in_english(both, case):
    chinese, english = both(case)
    assert english == chinese


def test_the_english_tour_sample_names_its_one_question_in_english(monkeypatch):
    use(monkeypatch, "en")
    from bridgeflow.api.main import app
    with TestClient(app) as client:
        summary = client.post("/batches/demo").json()
        master = client.get(f"/integration/batches/{summary['batch_id']}").json()
        sources = client.get(f"/batches/{summary['batch_id']}/sources").json()
    assert [(i["kind"], i["field"]) for i in master["issues"]] == [("disagreement", "Customer name")]
    names = sorted(s["filename"] for s in sources["sources"])
    assert names == ["Sample-Finance-2024-07.xlsx", "Sample-Marketing-2024-07.xlsx",
                     "Sample-Procurement-2024-07.xlsx", "Sample-Production-2024-07.xlsx"]


def test_the_english_renamed_column_is_asked_about_by_its_declared_name(monkeypatch):
    use(monkeypatch, "en")
    from bridgeflow.api.main import app
    with TestClient(app) as client:
        batch = client.post("/batches/demo?case=other").json()
        inbox = client.get(f"/monthly/inbox?period={batch['period']}").json()
    assert "marketing: cumulative_collections is missing or unreadable" in " | ".join(batch["review_blockers"])
    assert [q["subject"] for q in inbox["items"] if q["kind"] == "column_question"] == ["cumulative_collections"]


def test_the_english_discovery_sample_scopes_the_english_workflow_sample(monkeypatch):
    # The same chain as tests/test_workflow_scope.py, on the English project, policies and catalogue.
    from conftest import receipt

    use(monkeypatch, "en")
    from bridgeflow.api.main import app
    with TestClient(app) as client:
        decision = client.post("/discovery/sample")
        assert decision.status_code == 200, decision.text
        assert decision.json()["status"] == "approved"
        offered = client.get("/workflow/scope").json()["pending"]
        body = json.dumps({"project_id": "demo-handoff", "decision_id": "mvp",
                           "decision_seq": offered[0]["decision_seq"]}, separators=(",", ":")).encode()
        accepted = client.post("/tools/workflow-accept-scope", content=body,
                               headers={"content-type": "application/json", "x-bridgeflow-approval": receipt(body)})
        assert accepted.status_code == 200, accepted.text
        board = client.post("/workflow/sample")
        assert board.status_code == 200, board.text
        assert client.get("/batches/demo/cases").json()["cases"][0]["title"][1] == "Guided sample: one cross-department mismatch"
    text = json.dumps(board.json(), ensure_ascii=False)
    assert "Production record" in text and not CJK.search(text)


def test_the_english_set_is_what_its_generator_makes_and_carries_no_chinese():
    checked = subprocess.run([sys.executable, str(REPO_ROOT / "scripts/make_english_samples.py"), "--check"],
                             capture_output=True, text=True, check=False)
    assert checked.returncode == 0, checked.stderr
    for path in (REPO_ROOT / "data/demo_en").rglob("*.yaml"):
        if path.name != "glossary.yaml":
            assert not CJK.search(path.read_text(encoding="utf-8")), path


def test_the_english_tour_case_matches_its_manifest():
    manifest = json.loads((REPO_ROOT / "data/demo_en/demo/manifest.json").read_text(encoding="utf-8"))
    for relative, expected in manifest["files"].items():
        assert hashlib.sha256((REPO_ROOT / relative).read_bytes()).hexdigest() == expected, relative
    assert manifest["expected"]["open_questions"] == [{"kind": "disagreement", "field": "Customer name"}]


def test_the_guest_instance_loads_the_english_set_unless_told_otherwise(tmp_path, monkeypatch):
    sys.path.insert(0, str(REPO_ROOT / "scripts"))
    import start_web
    english = start_web.guest_sample_set({})
    assert english["demo_cases"] == REPO_ROOT / "data/demo_en/cases/cases.yaml"
    chinese = start_web.guest_sample_set({"BRIDGEFLOW_GUEST_SAMPLE_SET": SETS["zh"]})
    assert chinese["integration_spec"] == REPO_ROOT / "data/company_templates/integration.yaml"
    broken = tmp_path / "set.yaml"
    broken.write_text("dictionary: data/nope.yaml\n", encoding="utf-8")
    with pytest.raises(SystemExit):
        start_web.guest_sample_set({"BRIDGEFLOW_GUEST_SAMPLE_SET": str(broken)})


def test_the_english_downloads_read_in_english(monkeypatch, tmp_path):
    # The Word report and the master workbook take the interface's language; with the English
    # set behind them, nothing a judge downloads carries Chinese.
    import base64
    import io
    import zipfile

    import openpyxl
    from test_conclusions import finalize

    use(monkeypatch, "en")
    monkeypatch.setattr(settings, "result_store_path", str(tmp_path))
    from bridgeflow.api.main import app
    with TestClient(app) as client:
        batch = client.post("/batches/demo").json()["batch_id"]
        finalize(client, batch)
        report = client.get(f"/conclusions/batches/{batch}/report?lang=en").json()
        workbook = client.get(f"/integration/batches/{batch}/xlsx?lang=en").json()
    assert report["filename"].startswith("Monthly-conclusions-2024-07-")
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(report["base64"]))) as archive:
        text = archive.read("word/document.xml").decode()
    assert "8. Appendix: source index" in text and "VAT rate" in text and "Unconfirmed" in text
    assert not CJK.search(text), text[max(0, CJK.search(text).start() - 200):CJK.search(text).start() + 60]
    assert workbook["filename"].startswith("Cross-department-master-")
    book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(workbook["base64"])))
    assert book.sheetnames == ["Master", "Open items", "Conventions"]
    assert not any(isinstance(v, str) and CJK.search(v)
                   for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for v in row)
