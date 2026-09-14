"""Cross-path regressions: frozen declarations, original locations and safe outputs."""
import base64
import io
import json

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient

from bridgeflow import integration
from bridgeflow.api.batches import batch_path
from bridgeflow.api.main import app
from bridgeflow.config import REPO_ROOT, settings

BASE = REPO_ROOT / 'data/company_templates'


@pytest.fixture
def client(monkeypatch, tmp_path):
    spec = tmp_path / 'integration.yaml'
    spec.write_bytes((BASE / 'integration.yaml').read_bytes())
    monkeypatch.setattr(settings, 'integration_spec_path', str(spec))
    monkeypatch.setattr(settings, 'field_dictionary_path', str(REPO_ROOT / 'data/business_demo/dictionary.yaml'))
    with TestClient(app) as value:
        yield value


def upload(client, *, title=False, edit=None):
    files = []
    for department in integration.load_spec().departments:
        book = openpyxl.load_workbook(BASE / 'example' / f'{department}.xlsx')
        if edit:
            edit(department, book.active)
        if title:
            book.active.insert_rows(1)
            book.active.cell(1, 1, 'Synthetic title')
        buffer = io.BytesIO()
        book.save(buffer)
        files.append(('files', (f'{department}.xlsx', buffer.getvalue())))
    data = {'period': '2024-05', 'departments': list(integration.load_spec().departments)}
    if title:
        data['header_rows'] = ['2'] * len(files)
    response = client.post('/batches', data=data, files=files)
    assert response.status_code == 200, response.text
    return response.json()['batch_id']


def test_old_master_and_download_do_not_change_when_declaration_changes(client):
    batch = upload(client)
    url = f'/integration/batches/{batch}'
    original = client.get(url).json()
    path = integration.spec_path()
    raw = yaml.safe_load(path.read_text())
    raw['version'] = 'changed-after-import'
    raw['constants']['增值税税率'] = 0.03
    path.write_text(yaml.safe_dump(raw, allow_unicode=True))
    assert client.get(url).json() == original
    book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(client.get(url + '/xlsx').json()['base64'])))
    assert [c.value for c in book['总表'][2]] == list(original['rows'][0]['values'].values())
    new = upload(client)
    assert client.get(f'/integration/batches/{new}').json()['version'] == 'changed-after-import'


def test_legacy_batch_cannot_silently_borrow_current_integration_policy(client):
    batch = upload(client)
    path = batch_path(batch)
    raw = json.loads(path.read_text())
    raw.pop('integration_snapshot', None)
    path.write_text(json.dumps(raw))
    response = client.get(f'/integration/batches/{batch}')
    assert response.status_code == 409


def test_selected_header_keeps_original_locations_in_preview_and_master(client):
    batch = upload(client, title=True)
    preview = client.get(f'/batches/{batch}/sources/production').json()
    assert preview['header_row'] == 2
    assert preview['row_numbers'] == [3]
    master = client.get(f'/integration/batches/{batch}').json()
    assert master['departments_read']['production']['header_row'] == 2
    assert master['rows'][0]['provenance']['生产_生产量']['row'] == 3
    snapshot = json.loads(batch_path(batch).read_text())
    assert all(n >= 3 for table in snapshot['clean_tables'] for n in table['source_rows'])


def test_tool_summary_does_not_forward_numeric_conflict_values(client):
    def edit(department, ws):
        if department == 'procurement':
            for cell in ws[1]:
                if cell.value == '单方不含税毛利':
                    ws.cell(2, cell.column, 987654321)
    batch = upload(client, edit=edit)
    master = client.get(f'/integration/batches/{batch}').json()
    assert '987654321' in json.dumps(master)
    response = client.post('/tools/integration-summary', json={'batch_id': batch})
    assert response.status_code == 200
    assert response.json()['issues_by_kind']['derived_mismatch'] >= 1
    assert '987654321' not in response.text
    assert 'procurement.xlsx' not in response.text


def example_sheets():
    spec = integration.load_spec(BASE / 'integration.yaml')
    return spec, [integration.read_sheet(d, f'{d}.xlsx', (BASE / 'example' / f'{d}.xlsx').read_bytes())
                  for d in spec.departments]


@pytest.mark.parametrize('value', ['NaN', 'Infinity', '-Infinity'])
def test_nonfinite_numeric_cells_are_reported_not_crashed(value):
    spec, sheets = example_sheets()
    production = next(s for s in sheets if s.department == 'production')
    production.rows[0][production.headers.index('生产量')] = value
    result = integration.integrate(spec, sheets)
    assert any(i.kind == 'invalid_number' for i in result.issues)
    assert not result.rows[0].complete


def test_unverified_formula_never_counts_as_a_complete_row():
    spec, sheets = example_sheets()
    spec.constants['增值税税率'] = None
    result = integration.integrate(spec, sheets)
    assert any(i.kind == 'undeclared_constant' for i in result.issues)
    assert not result.rows[0].complete


def test_rollup_does_not_sum_only_the_rows_with_values():
    spec, sheets = example_sheets()
    production = next(s for s in sheets if s.department == 'production')
    production.rows.append(list(production.rows[0]))
    production.row_numbers.append(3)
    production.rows[1][production.headers.index('生产量')] = None
    result = integration.integrate(spec, sheets)
    assert not result.rows[0].complete
    assert any(i.kind == 'invalid_number' and i.field == '生产_生产量' for i in result.issues)
    assert result.rows[0].values['生产_生产量'] is None


def test_download_preserves_untrusted_formula_shaped_text_as_text():
    spec, sheets = example_sheets()
    result = integration.integrate(spec, sheets)
    name = next(iter(result.rows[0].values))
    result.rows[0].values[name] = '=1+1'
    result.issues.append(integration.Issue(kind='missing_key', field='=1+1', message='=1+1'))
    book = openpyxl.load_workbook(io.BytesIO(integration.to_xlsx(result)), data_only=False)
    assert book['总表'].cell(2, 1).value == '=1+1'
    assert book['总表'].cell(2, 1).data_type == 's'
    assert book['待确认'].cell(2, 5).data_type == 's'
