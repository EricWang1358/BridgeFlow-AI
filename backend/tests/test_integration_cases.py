"""The tuning sets and the holdout generator for the master table (#141).

The committed tuning sets (seeds 1 and 2) must grade clean and regenerate identically; a
fresh holdout set, from a seed nobody recorded, must grade clean too.
"""
import importlib.util
import json

import openpyxl

from bridgeflow.config import REPO_ROOT

_spec = importlib.util.spec_from_file_location("integration_cases", REPO_ROOT / "scripts" / "integration_cases.py")
cases = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(cases)


def test_the_committed_tuning_sets_grade_clean():
    for seed in cases.TUNING_SEEDS:
        report = cases.grade(cases.TUNING / f"seed-{seed}")
        assert report["passed"], report
        assert report["clean_keys"] >= 5 and report["split_keys"] >= 1
        assert {k for k in report["faults"]} >= {"derived_mismatch", "missing_department", "disagreement",
                                                  "check_failed", "invalid_number", "needs_rollup", "missing_key"}


def test_the_tuning_sets_are_what_their_seeds_generate(tmp_path):
    for seed in cases.TUNING_SEEDS:
        answer = cases.write(seed, tmp_path / str(seed))
        committed = cases.TUNING / f"seed-{seed}"
        assert answer == json.loads((committed / "answer.json").read_text(encoding="utf-8"))
        for path in committed.glob("*.xlsx"):
            fresh = openpyxl.load_workbook(tmp_path / str(seed) / path.name).active
            original = openpyxl.load_workbook(path).active
            assert [[c.value for c in r] for r in fresh.iter_rows()] == [[c.value for c in r] for r in original.iter_rows()]


def test_a_fresh_holdout_set_grades_clean():
    totals = cases.holdout(2)
    assert totals["passed_sets"] == 2, totals


def test_real_exports_are_graded_against_the_business_master_without_printing_values(tmp_path):
    import shutil
    for path in (cases.BASE / "example").glob("*.xlsx"):
        label = cases.integration.load_spec(cases.BASE / "integration.yaml").departments[path.stem].label
        shutil.copy(path, tmp_path / f"{label}(1).xlsx")  # named the way the business side names files
    shutil.copy(cases.BASE / "source" / "master-v2.xlsx", tmp_path / "expected.xlsx")
    report = cases.real(tmp_path)
    assert (report["expected_rows"], report["matched_rows"], report["exact_rows"], report["complete_rows"]) == (1, 1, 1, 1)
    assert all(report["department_files"].values()) and report["columns_with_differences"] == {}
    assert "项目A" not in json.dumps(report, ensure_ascii=False)

    book = openpyxl.load_workbook(tmp_path / "expected.xlsx")
    headers = [c.value for c in book.active[1]]
    book.active.cell(2, headers.index("市场_缺口") + 1, 1)
    book.save(tmp_path / "expected.xlsx")
    assert cases.real(tmp_path)["columns_with_differences"] == {"市场_缺口": 1}


def test_the_mock_monthly_exports_grade_as_their_readme_says():
    folder = REPO_ROOT / "data/mock_business/monthly"
    clean = cases.real(folder / "2024-05-调优A")
    assert (clean["matched_rows"], clean["complete_rows"], clean["exact_rows"], clean["issues_by_kind"]) == (4, 4, 4, {})
    june = cases.real(folder / "2024-06-调优B")
    assert june["issues_by_kind"].get("missing_department") == 1 and june["exact_rows"] == 3
    july = cases.real(folder / "2024-07-模拟留出")
    assert {"disagreement", "check_failed", "invalid_number"} <= set(july["issues_by_kind"])
    assert set(july["columns_with_differences"]) == {"市场_可争取", "物资_当月生产量", "物资_当月物资总成本"}
