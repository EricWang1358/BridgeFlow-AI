"""Correcting one department's file as a new batch version (E14-UC04).

The demo case's production file is spoiled here in exactly the way the acceptance describes
(a customer written as an abbreviation), imported, then corrected by re-uploading the real
file. Nothing calls a model.
"""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from test_conclusions import finalize

from bridgeflow.api.batches import batch_path, load_batch
from bridgeflow.api.main import app
from bridgeflow.config import REPO_ROOT, settings

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DEMO = REPO_ROOT / "data/mock_business/demo"
JUNE = REPO_ROOT / "data/mock_business/monthly/2024-06-调优B"
LABELS = {"production": "生产部", "procurement": "物资部", "finance": "财务部", "marketing": "市场部"}


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "field_dictionary_path", str(DEMO / "dictionary.yaml"))
    monkeypatch.setattr(settings, "result_store_path", str(tmp_path))
    with TestClient(app) as test_client:
        yield test_client


def abbreviated(payload: bytes) -> bytes:
    """The production file as a department would spoil it: the customer's short name."""
    book = openpyxl.load_workbook(io.BytesIO(payload))
    sheet = book.active
    column = [cell.value for cell in sheet[1]].index("客户单位") + 1
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, column).value:
            sheet.cell(row, column).value = str(sheet.cell(row, column).value).replace("示例建工第一分公司", "示例建工一分")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def original(name: str) -> bytes:
    return (DEMO / f"{name}.xlsx").read_bytes()


def import_demo(client, spoil=True):
    files = []
    for name in LABELS:
        payload = abbreviated(original(name)) if spoil and name == "production" else original(name)
        files.append(("files", (f"{name}.xlsx", payload, XLSX)))
    response = client.post("/batches", data={"period": "2024-07", "departments": list(LABELS)}, files=files)
    assert response.status_code == 200, response.text
    return response.json()["batch_id"]


def resupply(client, batch_id, department="production", payload=None, period="2024-07", reason="客户单位写成了简称"):
    return client.post(f"/batches/{batch_id}/departments/{department}",
                       data={"period": period, "reason": reason},
                       files={"file": (f"{department}.xlsx", payload if payload is not None else original(department), XLSX)})


def digests(client, batch_id):
    return {s["id"]: s.get("sha256") for s in client.get(f"/batches/{batch_id}/sources").json()["sources"]}


def disagreements(client, batch_id):
    master = client.get(f"/integration/batches/{batch_id}").json()
    return [i for i in master["issues"] if i["kind"] == "disagreement"]


def test_correcting_one_department_clears_the_disagreement_and_leaves_the_others_untouched(client):
    batch = import_demo(client)
    spoiled = disagreements(client, batch)
    # The demo case has one declared disagreement of its own; the abbreviation adds another.
    clean_batch = import_demo(client, spoil=False)
    assert len(spoiled) == len(disagreements(client, clean_batch)) + 1
    before = digests(client, batch)
    response = resupply(client, batch)
    assert response.status_code == 200, response.text
    body = response.json()
    new_id = body["batch"]["batch_id"]
    assert new_id != batch and body["batch"]["derived_from"] == batch and body["replaced"] == "production"
    assert len(disagreements(client, new_id)) == len(spoiled) - 1
    after = digests(client, new_id)
    assert {d: after[d] for d in ("procurement", "finance", "marketing")} == {d: before[d] for d in ("procurement", "finance", "marketing")}
    assert after["production"] != before["production"]
    assert body["diff"]["issues_delta"].get("disagreement", 0) < 0
    assert body["diff"]["cells_changed"] > 0 and "客户名称" in body["diff"]["fields_changed"]


def test_the_original_batch_its_master_and_its_report_are_unchanged(client):
    batch = import_demo(client)
    report = finalize(client, batch)
    before_file = batch_path(batch).read_bytes()
    before_master = client.get(f"/integration/batches/{batch}").json()
    new_id = resupply(client, batch).json()["batch"]["batch_id"]
    assert batch_path(batch).read_bytes() == before_file
    assert client.get(f"/integration/batches/{batch}").json() == before_master
    assert client.get(f"/batches/{batch}/review").json()["report_id"] == report["report_id"]
    # The version chain is readable from the old batch although that batch was never rewritten.
    assert client.get(f"/batches/{batch}").json()["superseded_by"] == [new_id]
    assert load_batch(batch).replacements == [] and len(load_batch(new_id).replacements) == 1
    assert load_batch(new_id).replacements[0]["reason"] == "客户单位写成了简称"
    # The derived batch carries no report of its own: the review must be run again.
    assert client.get(f"/batches/{new_id}/review").status_code == 404


def test_an_identical_file_derives_nothing(client):
    batch = import_demo(client, spoil=False)
    response = resupply(client, batch)
    assert response.status_code == 409 and "identical" in response.text
    assert client.get(f"/batches/{batch}").json()["superseded_by"] == []


def test_a_file_from_another_month_is_refused_by_its_own_rows(client):
    batch = import_demo(client)
    june = (JUNE / "生产部.xlsx").read_bytes()
    # The form says July, as an uploader who grabbed the wrong file would: the refusal comes
    # from the rows, not from the uploader's own claim.
    response = resupply(client, batch, payload=june)
    assert response.status_code == 422 and "2024-06" in response.text
    assert resupply(client, batch, period="2024-06").status_code == 422


def test_a_department_the_batch_does_not_have_is_refused(client):
    files = [("files", ("production.xlsx", original("production"), XLSX))]
    batch = client.post("/batches", data={"period": "2024-07", "departments": ["production"]}, files=files).json()["batch_id"]
    assert resupply(client, batch, department="finance", payload=original("finance")).status_code == 404
