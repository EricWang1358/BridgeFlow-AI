"""Which worksheet and which header row an XLSX import reads: declared or chosen, never guessed (#47).

The layout case is real: the business side's v1 物资部 template has a merged title row above
its header. Row 1 is not silently taken as the header, and the first sheet of several is not
silently taken as the table.
"""
import io

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient

from bridgeflow.api.main import app
from bridgeflow.config import REPO_ROOT, settings

TITLED = (REPO_ROOT / "data/company_templates/source/procurement-v1.xlsx").read_bytes()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def titled_with_a_row() -> bytes:
    book = openpyxl.load_workbook(io.BytesIO(TITLED))
    sheet = book.active
    width = sheet.max_column
    sheet.append([f"v{i}" for i in range(width)])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def workbook(*sheets, formula=False) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in sheets:
        sheet = book.create_sheet(name)
        sheet.append(["sku", "qty", "total"])
        sheet.append(["SKU-A1", 2, "=B2*3" if formula else 6])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    raw = yaml.safe_load((REPO_ROOT / "data/business_demo/dictionary.yaml").read_text(encoding="utf-8"))
    path = tmp_path / "dictionary.yaml"

    def declare(layout):
        path.write_text(yaml.safe_dump({**raw, "sheet_layout": layout} if layout else raw, allow_unicode=True), encoding="utf-8")
        monkeypatch.setattr(settings, "field_dictionary_path", str(path))
    declare(None)
    with TestClient(app) as test_client:
        test_client.declare = declare
        yield test_client


def upload(client, content, name="procurement.xlsx", **choice):
    return client.post("/batches", data={"period": "2025-11", "departments": "procurement", **choice},
                       files={"files": (name, content, XLSX)})


def test_a_title_row_above_the_header_is_refused_with_the_likely_header_row(client):
    response = upload(client, titled_with_a_row())
    assert response.status_code == 422
    assert "row 1 is not a header" in response.text and "row 2" in response.text


def test_a_chosen_header_row_reads_the_real_header(client):
    response = upload(client, titled_with_a_row(), header_rows="2")
    assert response.status_code == 200, response.text
    batch_id = response.json()["batch_id"]
    source = client.get(f"/batches/{batch_id}/sources/procurement").json()
    header = [c.value for c in openpyxl.load_workbook(io.BytesIO(TITLED)).active[2]]
    assert source["columns"][:3] == [str(h) for h in header[:3]]


def test_the_dictionary_can_declare_the_layout_once(client):
    client.declare({"procurement": {"header_row": 2}})
    assert upload(client, titled_with_a_row()).status_code == 200
    client.declare({"procurement": {"header_row": 0}})
    assert upload(client, titled_with_a_row()).status_code == 503


def test_several_sheets_are_refused_with_their_names_until_one_is_chosen(client):
    content = workbook("说明", "明细")
    refused = upload(client, content)
    assert refused.status_code == 422 and "说明" in refused.text and "明细" in refused.text
    assert upload(client, content, sheets="明细").status_code == 200
    assert "no worksheet named" in upload(client, content, sheets="汇总").text


def test_formulas_without_saved_results_are_refused_not_read_as_blanks(client):
    refused = upload(client, workbook("Sheet1", formula=True))
    assert refused.status_code == 422 and "formulas without saved results" in refused.text
    assert upload(client, workbook("Sheet1")).status_code == 200


def test_choices_must_line_up_with_the_files(client):
    response = client.post("/batches", data={"period": "2025-11", "departments": "procurement", "header_rows": ["1", "2"]},
                           files={"files": ("procurement.xlsx", workbook("Sheet1"), XLSX)})
    assert response.status_code == 422
