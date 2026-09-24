"""The department template with last month's values carried over (E14-UC02)."""
import base64
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from test_comparison import JUNE, LABELS, XLSX

from bridgeflow.api.main import app
from bridgeflow.config import REPO_ROOT, settings

COMPANY = REPO_ROOT / "data/company_templates/source"


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "field_dictionary_path", str(REPO_ROOT / "data/mock_business/demo/dictionary.yaml"))
    monkeypatch.setattr(settings, "result_store_path", str(tmp_path))
    with TestClient(app) as test_client:
        yield test_client


def import_june(client):
    files = [("files", (f"{label}.xlsx", (JUNE / f"{label}.xlsx").read_bytes(), XLSX)) for label in LABELS.values()]
    response = client.post("/batches", data={"period": "2024-06", "departments": list(LABELS)}, files=files)
    assert response.status_code == 200, response.text
    return response.json()["batch_id"]


def template(client, department="production", period="2024-07"):
    response = client.get(f"/batches/templates/{department}?period={period}")
    assert response.status_code == 200, response.text
    body = response.json()
    book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(body["base64"])))
    return body, book


def test_the_header_is_the_approved_templates_own_column_for_column(client):
    body, book = template(client)
    approved = openpyxl.load_workbook(COMPANY / "production-v2.xlsx").active
    assert [c.value for c in book.worksheets[0][1]] == [c.value for c in approved[1]]
    assert body["declaration"] and body["declaration"] in body["filename"]
    assert "填写说明" in book.sheetnames


def test_last_months_value_is_prefilled_and_says_where_it_came_from(client):
    batch = import_june(client)
    body, book = template(client)
    assert body["prefilled_from"] == "2024-06"
    sheet = book.worksheets[0]
    header = [c.value for c in sheet[1]]
    carried, project = header.index("上月实际量") + 1, header.index("项目编号") + 1
    rows = {sheet.cell(r, project).value: sheet.cell(r, carried) for r in range(2, sheet.max_row + 1)}
    master = client.get(f"/integration/batches/{batch}").json()
    expected = {row["key"][0]: row["values"]["生产_实际量"] for row in master["rows"]
                if row["values"].get("生产_实际量") is not None}
    assert rows and set(rows) <= set(expected)
    for key, cell in rows.items():
        assert cell.value == expected[key]
        assert "预填自 2024-06" in (cell.comment.text if cell.comment else ""), key
    notes = "\n".join(str(c.value) for row in book["填写说明"].iter_rows() for c in row if c.value)
    assert "2024-06" in notes and "上月实际量" in notes and "生产_实际量" in notes


def test_without_a_prior_batch_nothing_is_prefilled_and_the_notes_say_why(client):
    body, book = template(client)
    assert body["prefilled_from"] == "" and "2024-06" in body["reason"]
    sheet = book.worksheets[0]
    assert sheet.max_row == 1  # header only
    notes = "\n".join(str(c.value) for row in book["填写说明"].iter_rows() for c in row if c.value)
    assert "没有 2024-06 的批次" in notes


def test_a_real_department_template_does_not_carry_values_from_sample_history(client):
    client.post("/batches/demo/history")
    body, book = template(client)
    assert body["prefilled_from"] == ""
    assert book.worksheets[0].max_row == 1


def test_a_department_without_a_carry_over_declaration_is_not_prefilled_but_still_downloadable(client):
    import_june(client)
    body, book = template(client, department="finance")
    assert body["prefilled_from"] == "" and "没有声明" in body["reason"]
    approved = openpyxl.load_workbook(COMPANY / "finance-v2.xlsx").active
    # The approved template's own empty rows are handed on untouched: it is their template.
    assert book.worksheets[0].max_row == approved.max_row
    assert all(cell.value is None for row in book.worksheets[0].iter_rows(min_row=2) for cell in row)
    assert [c.value for c in book.worksheets[0][1]] == [c.value for c in approved[1]]


def test_an_undeclared_department_is_refused(client):
    assert client.get("/batches/templates/logistics?period=2024-07").status_code == 404


def test_the_prefilled_workbook_imports_under_the_same_declaration(client):
    """A template that cannot be handed back is not a template (E14-UC03 shares the check chain)."""
    import_june(client)
    _body, book = template(client)
    buffer = io.BytesIO()
    book.save(buffer)
    response = client.post("/batches/self-check", data={"period": "2024-07", "department": "production"},
                           files={"file": ("production.xlsx", buffer.getvalue(), XLSX)})
    assert response.status_code == 200, response.text
    # It is not a finished submission — the month's own figures are still empty — but the
    # check chain must be able to read it rather than refuse the file outright.
    assert response.json()["filename"] == "production.xlsx"
