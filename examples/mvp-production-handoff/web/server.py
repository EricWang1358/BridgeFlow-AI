"""Local workflow MVP v0.2. Demo roles, deterministic rules, simulated notifications."""
import argparse
import base64
import csv
import hashlib
import io
import json
import os
import re
import secrets
import sqlite3
import threading
import uuid
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FIELDS = ['报表月','厂站','年份','日期','客户单位','项目名称','货品名称','生产量','出厂量','实际量','备注']
ALIASES = dict(zip(['月份','站点','年','开单日期','客户','工程','产品','产量','发出量','确认数量','说明'], FIELDS))
SAMPLE = {'月份':'2026/09','站点':' 演示厂站A ','年':'2026','开单日期':'2026/09/01','客户':'演示客户A','工程':'演示项目A','产品':'演示产品A','产量':'100 方','发出量':'98 m³','说明':'合成演示；请补实际量及依据'}
DEFS = ['本案例的归属月份，须与业务日期一致','合成业务发生厂站','业务年份','业务发生日期','客户显示名称，不作为唯一键','项目显示名称，不作为唯一键','产品显示名称，规格待正式字典确认','独立记录的生产数量','独立记录的出厂数量','员工确认的实际业务量，不假定等于签收量','可选说明']
TEMPLATE = {'version':'production-pilot-v0.2','fields':[{'name':n,'definition':DEFS[i],'required':n!='备注','unit':'m³（演示）' if n in FIELDS[7:10] else '', 'aliases':[k for k,v in ALIASES.items() if v==n]} for i,n in enumerate(FIELDS)],'scope':'单条生产记录 → 市场部核对；仅模拟 MVP 决策与一线模板批准，不是真实投票/试点结果。'}

def now(): return datetime.now(timezone.utc).isoformat()
def dump(x): return json.dumps(x, ensure_ascii=False)

class Problem(Exception):
    def __init__(self, code, message): self.code,self.message=code,message


def normalize(raw):
    if not isinstance(raw,dict): raise Problem(422,'输入必须是一条字段记录')
    values, sources, errors = {}, {}, []
    for source,value in raw.items():
        name=ALIASES.get(source,source)
        if name not in FIELDS:
            errors.append('未知字段：'+source); continue
        if name in values:
            errors.append('字段重复映射：'+name); continue
        if isinstance(value,(list,dict,bool)): errors.append(name+'不是可接受的单值'); continue
        values[name]=str(value).strip() if value is not None else ''
        sources[name]={'source_field':source,'original_value':value}
    for name in FIELDS[7:10]:
        s=values.get(name,'')
        if s:
            match=re.fullmatch(r'([+]?(?:\d+(?:\.\d*)?|\.\d+))\s*(?:方|m³)?',s)
            if not match:
                errors.append(name+'须为非负数，单位仅支持本例的方或 m³'); continue
            d=Decimal(match[1])
            if d>Decimal('1000000000'): errors.append(name+'超过演示上限'); continue
            values[name]=format(d,'f')
    for name,fmt in [('日期','%Y-%m-%d'),('报表月','%Y-%m')]:
        if values.get(name):
            try: values[name]=datetime.strptime(values[name].replace('/','-'),fmt).strftime(fmt)
            except ValueError: errors.append(name+'格式无法解析')
    if values.get('日期') and len(values['日期'])==10:
        if values.get('报表月') and values['日期'][:7]!=values['报表月']: errors.append('日期与报表月不一致')
        if values.get('年份') and values['日期'][:4]!=values['年份']: errors.append('日期与年份不一致')
    values={name:values.get(name,'') for name in FIELDS}
    missing=[n for n in FIELDS[:-1] if values[n]=='']
    delta=None
    if not errors and all(values[n]!='' for n in ['出厂量','实际量']): delta=str(Decimal(values['出厂量'])-Decimal(values['实际量']))
    return {'values':values,'sources':sources,'missing':missing,'errors':errors,'valid':not errors and not missing,'difference':delta}


def parse_file(name, encoded):
    try: data=base64.b64decode(encoded,validate=True)
    except Exception: raise Problem(422,'文件编码无效')
    if len(data)>1024*1024: raise Problem(413,'本示例文件上限 1 MB')
    ext=Path(name).suffix.lower()
    if ext=='.json':
        try: obj=json.loads(data.decode('utf-8-sig'))
        except Exception: raise Problem(422,'无法解析 JSON')
        if isinstance(obj,list):
            if len(obj)!=1: raise Problem(422,'本 MVP 一次只处理一条业务记录')
            obj=obj[0]
        if not isinstance(obj,dict): raise Problem(422,'JSON 应是一条对象记录')
        return obj
    if ext=='.csv':
        try: rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        except Exception: raise Problem(422,'请上传 UTF-8 CSV')
    elif ext=='.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist())>10*1024*1024: raise Problem(413,'解压内容超过演示限制')
            from openpyxl import load_workbook
            w=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            try:
                if len(w.sheetnames)!=1: raise Problem(422,'本 MVP 仅支持单工作表')
                s=w.active
                if s.max_row>500 or s.max_column>40: raise Problem(422,'表格超出单记录演示范围')
                rows=[]
                for row in s.iter_rows():
                    if any(c.data_type=='f' for c in row): raise Problem(422,'请先核验公式；本示例不执行上传公式')
                    rows.append([c.value.isoformat()[:10] if isinstance(c.value,datetime) else c.value for c in row])
            finally: w.close()
        except ImportError: raise Problem(422,'XLSX 需要 openpyxl；请使用项目的 Python 环境')
        except Problem: raise
        except Exception: raise Problem(422,'无法读取 XLSX')
    else: raise Problem(422,'只支持 CSV、JSON、XLSX；本版本不支持图片/OCR')
    rows=[r for r in rows if any(v is not None and str(v).strip() for v in r)]
    if len(rows)<2: raise Problem(422,'只有表头，没有业务数据；请补一条记录或载入合成样例')
    if len(rows)!=2: raise Problem(422,'本 MVP 仅支持一行表头加一条业务记录')
    headers=[str(v).strip() if v is not None else '' for v in rows[0]]
    if any(not h for h in headers) or len(set(headers))!=len(headers): raise Problem(422,'表头有空白或重复项')
    if len(rows[1])!=len(headers): raise Problem(422,'数据列数与表头不一致')
    return dict(zip(headers,rows[1]))

class Store:
    def __init__(self,path):
        self.path=path; self.lock=threading.Lock()
        with self.connect() as c:
            c.executescript('CREATE TABLE IF NOT EXISTS docs(id TEXT PRIMARY KEY,body TEXT);CREATE TABLE IF NOT EXISTS submissions(id TEXT,version INT,body TEXT,PRIMARY KEY(id,version));CREATE TABLE IF NOT EXISTS outbox(id TEXT,version INT,state TEXT,attempts INT,PRIMARY KEY(id,version));')
    def connect(self): return sqlite3.connect(self.path,timeout=10)
    def event(self,d,event,note): d['events'].append({'time':now(),'event':event,'note':note})
    def save(self,c,d): c.execute('INSERT OR REPLACE INTO docs VALUES(?,?)',(d['id'],dump(d)))
    def all(self):
        with self.connect() as c: return [json.loads(r[0]) for r in c.execute('SELECT body FROM docs ORDER BY rowid DESC')]
    def create(self,raw,filename):
        checked=normalize(raw)
        d={'id':'P-'+uuid.uuid4().hex[:8],'version':1,'template_version':TEMPLATE['version'],'filename':filename,'raw':raw,'checked':checked,'state':'draft','approved_version':None,'confirmation':'','notification':'not_created','receipt':None,'events':[]}
        self.event(d,'收到材料',filename)
        self.event(d,'提取与校验','缺失：'+('、'.join(checked['missing']) or '无')+'；错误：'+('、'.join(checked['errors']) or '无'))
        with self.lock, self.connect() as c: self.save(c,d)
        return d
    def action(self,id,action,p):
        with self.lock, self.connect() as c:
            r=c.execute('SELECT body FROM docs WHERE id=?',(id,)).fetchone()
            if not r: raise Problem(404,'记录不存在')
            d=json.loads(r[0]); v=d['version']
            if p.get('version')!=v: raise Problem(409,'记录版本已变化，请刷新后重试')
            note=str(p.get('note','')).strip()
            if action=='patch':
                if d['state'] not in ['draft','returned']: raise Problem(409,'已提交的版本不可编辑；请先由市场部退回')
                if not note: raise Problem(422,'请填写补充或修改的依据')
                checked=normalize(p.get('values'))
                # Preserve uploaded source plus each manual edit in the timeline.
                changes={n:{'before':d['checked']['values'].get(n),'after':checked['values'][n]} for n in FIELDS if d['checked']['values'].get(n)!=checked['values'][n]}
                if not changes: raise Problem(422,'没有字段变化')
                for n in FIELDS:
                    checked['sources'][n]=({'source_field':'人工补充','original_value':checked['values'][n],'note':note} if n in changes else d['checked']['sources'].get(n,checked['sources'].get(n)))
                d.update(checked=checked,version=v+1,approved_version=None,confirmation='',state='draft',notification='not_created',receipt=None)
                self.event(d,'补充或修改',{'reason':note,'changes':changes})
            elif action=='approve':
                if d['state']!='draft' or not d['checked']['valid']: raise Problem(422,'请先补齐并校验全部必填字段')
                if not note: raise Problem(422,'请输入复核依据')
                d['approved_version']=v; d['confirmation']=note
                self.event(d,'生产部复核（演示角色）',note)
            elif action=='submit':
                old=c.execute('SELECT body FROM submissions WHERE id=? AND version=?',(id,v)).fetchone()
                if old: return d
                if d['state']!='draft' or d['approved_version']!=v or not normalize(d['checked']['values'])['valid']: raise Problem(422,'当前版本尚未完成有效复核')
                c.execute('INSERT INTO submissions VALUES(?,?,?)',(id,v,dump(d['checked']['values'])))
                c.execute('INSERT INTO outbox VALUES(?,?,?,?)',(id,v,'pending',0))
                d.update(state='pending_market',notification='pending',receipt={'id':id,'version':v,'target':'本地 SQLite / submissions','time':now()})
                self.event(d,'标准数据入库',d['receipt'])
                failed=p.get('simulate_notification_failure') is True
                d['notification']='failed' if failed else 'simulated_sent'
                c.execute('UPDATE outbox SET state=?,attempts=1 WHERE id=? AND version=?',(d['notification'],id,v))
                self.event(d,'模拟通知失败' if failed else '模拟通知已发送','生产部数据已就绪，待市场部处理；没有发送真实消息')
            elif action=='retry':
                if d['notification']!='failed': raise Problem(409,'当前没有失败通知需要重试')
                d['notification']='simulated_sent'
                c.execute("UPDATE outbox SET state='simulated_sent',attempts=attempts+1 WHERE id=? AND version=?",(id,v))
                self.event(d,'模拟通知重试成功','入库数据保持不变')
            elif action=='accept':
                if d['state']!='pending_market': raise Problem(409,'只能接收待市场部处理的记录')
                d['state']='market_processing'; self.event(d,'市场部接收（演示角色）','进入核对阶段')
            elif action=='return':
                if d['state'] not in ['pending_market','market_processing']: raise Problem(409,'当前阶段不能退回')
                if not note: raise Problem(422,'请填写退回原因')
                d.update(state='returned',approved_version=None)
                self.event(d,'市场部退回',note)
            elif action=='complete':
                if d['state']!='market_processing': raise Problem(409,'市场部须先接收再完成')
                if not note: raise Problem(422,'请填写核对结论；存在差异时说明处置')
                d['state']='completed'; self.event(d,'市场部完成（演示角色）',note)
            else: raise Problem(404,'动作不存在')
            self.save(c,d); return d


def make_server(db,port=3083):
    store=Store(db); token=secrets.token_urlsafe(24)
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,*args): pass
        def answer(self,code,data,kind='application/json; charset=utf-8'):
            body=(dump(data).encode() if kind.startswith('application/json') else data)
            self.send_response(code); self.send_header('Content-Type',kind); self.send_header('Cache-Control','no-store')
            self.send_header('X-Content-Type-Options','nosniff'); self.end_headers(); self.wfile.write(body)
        def safe_host(self): return self.headers.get('Host')==f'127.0.0.1:{self.server.server_port}'
        def do_GET(self):
            if not self.safe_host(): return self.answer(403,{'error':'仅允许本机地址'})
            if self.path=='/': return self.answer(200,(ROOT/'app.html').read_bytes(),'text/html; charset=utf-8')
            if self.path=='/api/state': return self.answer(200,{'token':token,'template':TEMPLATE,'docs':store.all(),'sample':SAMPLE})
            if self.path=='/sample.csv':
                out=io.StringIO(); w=csv.writer(out); w.writerow(SAMPLE.keys()); w.writerow(SAMPLE.values())
                return self.answer(200,('\ufeff'+out.getvalue()).encode(),'text/csv; charset=utf-8')
            if self.path.startswith('/export/'):
                id=self.path.split('/')[-1]; d=next((x for x in store.all() if x['id']==id),None)
                if not d or not d['receipt']: return self.answer(404,{'error':'没有已入库版本可导出'})
                out=io.StringIO(); w=csv.writer(out); w.writerow(FIELDS)
                # Neutralize spreadsheet formula injection on text export.
                vals=[d['checked']['values'][n] for n in FIELDS]
                w.writerow(["'"+s if s.lstrip().startswith(('=','+','-','@')) else s for s in vals])
                return self.answer(200,('\ufeff'+out.getvalue()).encode(),'text/csv; charset=utf-8')
            return self.answer(404,{'error':'不存在'})
        def do_POST(self):
            if not self.safe_host() or self.headers.get('X-Demo-Token')!=token: return self.answer(403,{'error':'请从本机页面操作'})
            origin=self.headers.get('Origin')
            if origin and origin!=f'http://127.0.0.1:{self.server.server_port}': return self.answer(403,{'error':'来源不允许'})
            try:
                size=int(self.headers.get('Content-Length','0'))
                if size<0 or size>1500000: raise Problem(413,'请求超过大小限制')
                p=json.loads(self.rfile.read(size))
                if not isinstance(p,dict): raise Problem(422,'请求须为对象')
                if self.path=='/api/create':
                    if p.get('sample'): raw,filename=SAMPLE.copy(),'内置合成样例'
                    else:
                        filename=Path(str(p.get('filename','upload'))).name
                        raw=parse_file(filename,p.get('data',''))
                    d=store.create(raw,filename)
                else:
                    parts=self.path.split('/')
                    if len(parts)!=5 or parts[1:3]!=['api','docs']: raise Problem(404,'接口不存在')
                    d=store.action(parts[3],parts[4],p)
                self.answer(200,{'doc':d})
            except Problem as e: self.answer(e.code,{'error':e.message})
            except (ValueError,TypeError,KeyError): self.answer(422,{'error':'请求格式不正确'})
    server=ThreadingHTTPServer(('127.0.0.1',port),Handler)
    return server,store

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--port',type=int,default=3083); parser.add_argument('--db',default=str(ROOT/'data'/'workflow.sqlite'))
    a=parser.parse_args(); Path(a.db).parent.mkdir(parents=True,exist_ok=True); os.umask(0o077)
    server,_=make_server(a.db,a.port)
    print(f'本地 MVP：http://127.0.0.1:{server.server_port} （规则演示，无模型或真实通知）',flush=True)
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close()
