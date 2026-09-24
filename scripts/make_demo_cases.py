"""Three demo cases for the same fictional concrete supplier, 2024-07 (#141 inputs).

One batch cannot show every state at once — a batch that imports cleanly cannot also be
held back — so the demo is three cases a person can switch between:

- ``core``   many problems at once: every error the held-out month plants (a short customer
             name, a quantity taken from the plan, text in a money field) plus a department
             that left a project out. Several kinds of open item appear together.
- ``other``  different problems from ``core``, none of them shared: a row from another
             month, a negative quantity, a declared column renamed in the sheet and a
             department template missing a required column.
- ``clean``  everything right: no planted error, and prices and collection plans set so
             every declared metric sits inside its threshold. It is ready to review.
- ``history/2024-05``, ``history/2024-06``  the same supplier's two earlier months with
             nothing planted, so the trend charts have more than one point.

All three reuse the simulation in ``make_mock_business.py`` so the books still tie, and the
dictionary of the retained sample (``data/mock_business/demo/dictionary.yaml``). The retained
tour sample itself is not touched; ``--check`` proves the simulation still reproduces it.

    python scripts/make_demo_cases.py           # writes data/mock_business/cases/
    python scripts/make_demo_cases.py --check   # the tour sample is still reproduced exactly
"""
from __future__ import annotations

import dataclasses
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_mock_business as mock  # noqa: E402

from bridgeflow.integration import load_spec  # noqa: E402

ROOT = mock.ROOT
OUT = ROOT / "data" / "mock_business" / "cases"
MONTH = mock.DEMO_MONTH

#: core: all of the held-out month's planted errors, and marketing leaves one project out.
CORE = mock.PLANTED[MONTH] + [("marketing", "PRJ2023098", "row", None, "市场部漏报了收尾项目")]


def clean_projects() -> list:
    """Prices 15% higher and collection plans that cover the contract debt: every metric in range."""
    return [dataclasses.replace(p, price=(p.price * Decimal("1.15")).quantize(Decimal("1"))) for p in mock.PROJECTS]


CLEAN_PLAN = ("0.60", "0.30", "0.15")


def _headers(template: Path) -> list[str]:
    return [str(c.value or "").strip() for c in openpyxl.load_workbook(template).active[1]]


def write(folder: Path, months: dict, planted: list, month: str = MONTH) -> None:
    spec = load_spec(mock.TEMPLATES / "integration.yaml")
    folder.mkdir(parents=True, exist_ok=True)
    for department, decl in spec.departments.items():
        mock._write_department(spec, department, month, months[month], _headers(mock.TEMPLATES / decl.template),
                               folder / f"{department}.xlsx", planted)


def plant_other(folder: Path) -> None:
    """Errors none of which the core case has, applied to otherwise correct files."""
    # production: a delivery row from June, and a negative produced quantity.
    book = openpyxl.load_workbook(folder / "production.xlsx")
    sheet = book.active
    header = [c.value for c in sheet[1]]
    june = [c.value for c in sheet[2]]
    june[header.index("报表月")], june[header.index("日期")] = 6, "2024-06-28"
    sheet.append(june)
    sheet.cell(row=3, column=header.index("生产量") + 1, value=-12)
    book.save(folder / "production.xlsx")
    # marketing: a declared column renamed in the sheet (累计收款 → 累计回款). An extra column
    # alone is harmless and ignored; a renamed one is a question the captain can propose on.
    book = openpyxl.load_workbook(folder / "marketing.xlsx")
    sheet = book.active
    header = [c.value for c in sheet[1]]
    sheet.cell(row=1, column=header.index("累计收款") + 1, value="累计回款")
    book.save(folder / "marketing.xlsx")
    # procurement: the template's 容重 column was deleted before sending.
    book = openpyxl.load_workbook(folder / "procurement.xlsx")
    sheet = book.active
    header = [c.value for c in sheet[1]]
    sheet.delete_cols(header.index("容重") + 1)
    book.save(folder / "procurement.xlsx")


def check() -> bool:
    """The parametrised simulation still reproduces the retained tour sample cell for cell."""
    spec = load_spec(mock.TEMPLATES / "integration.yaml")
    months = mock.simulate()
    same = True
    with tempfile.TemporaryDirectory() as scratch:
        for department, decl in spec.departments.items():
            path = Path(scratch) / f"{department}.xlsx"
            mock._write_department(spec, department, MONTH, months[MONTH], _headers(mock.TEMPLATES / decl.template),
                                   path, mock.DEMO_PLANTED)
            ours = [[c.value for c in r] for r in openpyxl.load_workbook(path).active.iter_rows()]
            kept = [[c.value for c in r] for r in openpyxl.load_workbook(ROOT / f"data/mock_business/demo/{department}.xlsx").active.iter_rows()]
            print(f"{department}: {'identical' if ours == kept else 'DIFFERENT'}")
            same &= ours == kept
    return same


def main() -> None:
    if "--check" in sys.argv:
        sys.exit(0 if check() else 1)
    default = mock.simulate()
    write(OUT / "core", default, CORE)
    write(OUT / "other", default, [])
    plant_other(OUT / "other")
    write(OUT / "clean", mock.simulate(clean_projects(), CLEAN_PLAN), [])
    # The same supplier's earlier months, nothing planted: history for the trends.
    for month, _label in mock.MONTHS:
        if month != MONTH:
            write(OUT / "history" / month, default, [], month)
    print(f"wrote {OUT.relative_to(ROOT)}/{{core,other,clean,history}}")


if __name__ == "__main__":
    main()
