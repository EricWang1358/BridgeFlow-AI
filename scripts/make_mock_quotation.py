"""模拟报价样板中的表格件（#104、#7、#20、#40）：成本与产能依据、标准报价单模板。

文字件（询价记录、合同范本、定价政策）直接写在 `data/mock_business/quotation/*.md`。这里生成：

- `04-成本与产能依据.xlsx`：物资部 C35 配合比材料成本（2024-07 价）、运输泵送及制造费用、
  西区拌站产能与在手订单、客户回款记录（财务部核实）；
- `03-标准报价单模板.xlsx`：人工报价单，每个数字旁写明取自哪份原件的哪个位置。

所有单位、项目与人名均为虚构。确定性输出，重跑得到同样的单元格。
"""
from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parents[1] / "data" / "mock_business" / "quotation"
D = Decimal
BOLD = Font(bold=True)

MIX_C35 = [  # 材料, 规格, 含税到厂价 元/吨, 进项税率, 含量 kg/方
    ("水泥", "P·O42.5", 405, "0.13", "320"),
    ("矿粉", "S95", 255, "0.13", "70"),
    ("粉煤灰", "F类Ⅱ级", 150, "0.13", "70"),
    ("机制砂", "细度模数 2.6–3.0", 88, "0.03", "770"),
    ("碎石", "5–25mm 连续级配", 78, "0.03", "1040"),
    ("聚羧酸减水剂", "固含量 20%", 2350, "0.13", "6.8"),
]

PAYMENTS = [  # 结算月, 结算日, 结算额, 回款日, 回款额
    ("2024-01", date(2024, 1, 25), 1215000.00, date(2024, 3, 22), 972000.00),
    ("2024-02", date(2024, 2, 26), 850500.00, date(2024, 4, 25), 680400.00),
    ("2024-03", date(2024, 3, 25), 1336500.00, date(2024, 5, 20), 1069200.00),
    ("2024-04", date(2024, 4, 25), 1458000.00, date(2024, 6, 24), 1166400.00),
    ("2024-05", date(2024, 5, 27), 1599750.00, date(2024, 7, 23), 1279800.00),
]


def q2(value: Decimal) -> Decimal:
    return value.quantize(D("0.01"), rounding=ROUND_HALF_UP)


def cost_basis() -> dict:
    book = openpyxl.Workbook()
    mix = book.active
    mix.title = "配合比与材料成本"
    mix.append(["C35 泵送混凝土配合比材料成本测算（2024-07 到厂价）", None, None, None, None, None, None])
    mix["A1"].font = BOLD
    mix.append(["材料", "规格", "含税单价(元/吨)", "进项税率", "含量(kg/方)", "含税成本(元/方)", "不含税成本(元/方)"])
    taxed_total = untaxed_total = D(0)
    for name, grade, price, rate, amount in MIX_C35:
        taxed = q2(D(price) * D(amount) / 1000)
        untaxed = q2(D(price) / (1 + D(rate)) * D(amount) / 1000)
        taxed_total += taxed
        untaxed_total += untaxed
        mix.append([name, grade, price, float(rate), float(amount), float(taxed), float(untaxed)])
    mix.append(["合计", None, None, None, None, float(taxed_total), float(untaxed_total)])
    material_row = mix.max_row
    mix.append(["测算：物资部 陈刚（模拟）　复核：物资部经理（模拟）　日期：2024-07-10"])

    fees = book.create_sheet("运输泵送与制造费用")
    fees.append(["项目", "元/方", "口径"])
    for row in [("运输（14 公里档）", 22.0, "罐车台班折算"), ("泵送", 16.0, "泵车台班折算"), ("制造费用", 12.0, "水电、人工、折旧分摊")]:
        fees.append(list(row))
    fees.append(["合计", 50.0, "定价政策第一条"])
    fee_row = fees.max_row

    capacity = book.create_sheet("产能与在手订单")
    capacity.append(["西区拌站 2024-08 产能测算", None, None])
    capacity["A1"].font = BOLD
    capacity.append(["项目", "方/月", "说明"])
    rows = [("设计月产能", 26000, "HZS180 两条线 × 有效台时"), ("保留急单产能 15%", -3900, "定价政策第四条"),
            ("在手订单占用", -18900, "PRJ2024017 等 5 个在供项目 8 月排产"), ("可承接方量", 3200, "= 设计 − 保留 − 占用")]
    for row in rows:
        capacity.append(list(row))
    capacity_row = capacity.max_row
    capacity.append(["新项目准备周期（天）", 7, "配合比试配 5 天 + 开盘鉴定 2 天"])
    lead_row = capacity.max_row
    capacity.append(["确认：生产部经理（模拟）　日期：2024-07-11"])

    history = book.create_sheet("客户回款记录")
    history.append(["客户 CUST0211 示例城建集团有限公司 · 项目 PRJ2024017 回款记录（财务部核实）"] + [None] * 5)
    history["A1"].font = BOLD
    history.append(["结算月", "结算日", "结算额(元)", "回款日", "回款额(元)", "回款天数"])
    days = []
    for month, settled_on, settled, paid_on, paid in PAYMENTS:
        days.append((paid_on - settled_on).days)
        history.append([month, settled_on, settled, paid_on, paid, days[-1]])
    average = q2(D(sum(days)) / len(days))
    history.append(["平均回款天数", None, None, None, None, float(average)])
    average_row = history.max_row
    history.append(["信用等级（定价政策第三条）", None, None, None, None, "B"])
    history.append(["资金占用成本（元/方）", None, None, None, None, 8.0])
    finance_row = history.max_row
    history.append(["核实：财务部 周洁（模拟）　日期：2024-07-11"])
    for sheet in book.worksheets:
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 18
    book.save(OUT / "04-成本与产能依据.xlsx")
    return {"material_cell": f"配合比与材料成本!G{material_row}", "material": untaxed_total,
            "fee_cell": f"运输泵送与制造费用!B{fee_row}", "capacity_cell": f"产能与在手订单!B{capacity_row}",
            "lead_cell": f"产能与在手订单!B{lead_row}", "days_cell": f"客户回款记录!F{average_row}", "days": average,
            "finance_cell": f"客户回款记录!F{finance_row}"}


def quote_template(cells: dict) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "报价单"
    sheet.append(["商品混凝土报价单（内部草稿，未经审批不得外发）"])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append(["报价编号", "BJ-2024-0715", "客户", "示例城建集团有限公司", "项目", "高新区产业园四号厂房"])
    sheet.append([])
    sheet.append(["项目", "数值", "单位", "取自原件", "位置", "填写人"])
    lines = [
        ("产品", "C35 泵送", "", "01-客户询价记录.md", "二、需求", "市场部"),
        ("预计总方量", None, "方", "01-客户询价记录.md", "二、需求", "市场部"),
        ("月峰值", None, "方/月", "01-客户询价记录.md", "二、需求", "市场部"),
        ("材料成本（不含税）", None, "元/方", "04-成本与产能依据.xlsx", cells["material_cell"], "物资部"),
        ("运输泵送及制造费用", None, "元/方", "04-成本与产能依据.xlsx", cells["fee_cell"], "物资部"),
        ("资金占用成本", None, "元/方", "04-成本与产能依据.xlsx", cells["finance_cell"], "财务部"),
        ("单方成本（不含税）", None, "元/方", "公式", "材料 + 运输泵送制造 + 资金", "系统计算"),
        ("含税底价", None, "元/方", "05-定价与信用政策.md", "二、价格带", "系统计算"),
        ("含税目标价", None, "元/方", "05-定价与信用政策.md", "二、价格带", "系统计算"),
        ("含税争取价", None, "元/方", "05-定价与信用政策.md", "二、价格带", "系统计算"),
        ("客户期望价", None, "元/方", "01-客户询价记录.md", "三、客户商务条件", "市场部"),
        ("付款比例（次月）", None, "", "01-客户询价记录.md", "三、客户商务条件", "市场部"),
        ("历史平均回款天数", None, "天", "04-成本与产能依据.xlsx", cells["days_cell"], "财务部"),
        ("可承接方量", None, "方/月", "04-成本与产能依据.xlsx", cells["capacity_cell"], "生产部"),
        ("准备周期", None, "天", "04-成本与产能依据.xlsx", cells["lead_cell"], "生产部"),
        ("报价有效期", 15, "天", "01-客户询价记录.md", "三、客户商务条件", "市场部"),
    ]
    for line in lines:
        sheet.append(list(line))
    sheet.append([])
    sheet.append(["审批（定价政策第五条）", "销售经理：", "销售总监：", "总经理：", "财务总监：", ""])
    sheet.append(["说明", "「数值」为空的行由报价系统按声明从原件抽取或计算后填入；人工只核对来源与审批。"])
    for column, width in zip("ABCDEF", (22, 16, 10, 26, 30, 12), strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    book.save(OUT / "03-标准报价单模板.xlsx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    cells = cost_basis()
    quote_template(cells)
    print(f"material cost {cells['material']} 元/方, average payment days {cells['days']}")


if __name__ == "__main__":
    main()
