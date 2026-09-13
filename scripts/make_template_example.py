"""Split the business side's worked example into the four v2 department templates.

The master template (`source/master-v2.xlsx`) carries one fully filled example row
(项目A / 客户A, synthetic). This writes the department sheets that row would have come
from, following `integration.yaml` column by column, so integration can be checked
against the business side's own answer. Deterministic; rerun to regenerate.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend" / "src"))
from bridgeflow.integration import load_spec  # noqa: E402

BASE = ROOT / "data" / "company_templates"


def main() -> None:
    spec = load_spec(BASE / "integration.yaml")
    master = openpyxl.load_workbook(BASE / spec.master_template).active
    example = dict(zip([c.value for c in master[1]], [c.value for c in master[2]], strict=True))
    out = BASE / "example"
    out.mkdir(exist_ok=True)
    for department, decl in spec.departments.items():
        template = openpyxl.load_workbook(BASE / decl.template).active
        headers = [c.value for c in template[1]]
        row = [None] * len(headers)
        for name, field in spec.fields.items():
            ref = field.sources.get(department)
            if ref is None or example.get(name) is None:
                continue
            if ref.period_from:
                year, month = str(example[name]).split("-")
                row[[str(h).strip() for h in headers].index(ref.period_from[0])] = int(year)
                row[[str(h).strip() for h in headers].index(ref.period_from[1])] = int(month)
                continue
            matches = [i for i, h in enumerate(headers) if str(h or "").strip() == ref.column]
            row[matches[ref.occurrence - 1]] = example[name]
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet.append(headers)
        sheet.append(row)
        book.save(out / f"{department}.xlsx")
    print(f"wrote {len(spec.departments)} department sheets to {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
