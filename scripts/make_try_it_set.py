"""自己上手试一遍用的材料：虚构商砼公司 2024-08 的四部门报表（#141 的延伸）。

和 `make_mock_business.py` 生成的调优 / 留出集不同，这一套是给人**手动走一遍**用的：
除了能正常导入的四份文件，还准备了会被系统拒绝的件，好让「拒绝什么、怎么说」能当场看到。

它复用同一套模拟：同样的四个项目、同样的配合比与回款习惯，只是多算一个月。
2024-08 的数字接在 07 之后（第二实验学校 8 月 15 日完工，所以收尾减量），账仍然互相勾稽。

    python scripts/make_try_it_set.py       # 写到 data/mock_business/try-it-2024-08/

所有单位、项目与人名均为虚构。
"""
from __future__ import annotations

import shutil
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "backend" / "src"))

import make_mock_business as gen  # noqa: E402
from bridgeflow.integration import load_spec  # noqa: E402

MONTH = "2024-08"
OUT = ROOT / "data" / "mock_business" / "try-it-2024-08"
D = Decimal

#: 8 月的设定：滨江与高新区继续，市政道路签收率继续下滑，第二实验学校 8 月中完工。
ACTUAL = {"PRJ2024011": 5700, "PRJ2024017": 5200, "PRJ2024023": 1950, "PRJ2023098": 600}
SIGN_RATE = {"PRJ2024011": "0.995", "PRJ2024017": "0.992", "PRJ2024023": "0.918", "PRJ2023098": "0.987"}
STATUS = {"PRJ2024011": "在供", "PRJ2024017": "在供", "PRJ2024023": "在供", "PRJ2023098": "收尾"}
PRICES = [410, 258, 150, 90, 78, 2380]

#: 手动试用时要看到的那处填报差异：生产部把客户单位写成简称，总表因此报「名称不一致」，
#: 而补传更正件里写的是全称——这一对正好演示 E14-UC04。
PLANTED = [("production", "PRJ2024017", "客户名称", "示例城建集团", "生产部客户单位写了简称")]


def extend_simulation() -> dict:
    """把 8 月接到已有三个月之后，返回该月的模拟结果。

    追加在 MONTHS 末尾，所以 05–07 的随机抽样顺序不变——这套脚本从不重写已有月份的文件。
    """
    gen.MONTHS.append((MONTH, "试用"))
    gen.MATERIAL_PRICE[MONTH] = PRICES
    for project in gen.PROJECTS:
        project.monthly_actual[MONTH] = ACTUAL[project.code]
        project.sign_rate[MONTH] = SIGN_RATE[project.code]
        project.status[MONTH] = STATUS[project.code]
    return gen.simulate()[MONTH]


def write_departments(spec, month_data: dict, folder: Path, planted: list) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for department, decl in spec.departments.items():
        template = openpyxl.load_workbook(gen.TEMPLATES / decl.template).active
        headers = [str(cell.value or "").strip() for cell in template[1]]
        gen._write_department(spec, department, MONTH, month_data, headers, folder / f"{decl.label}.xlsx", planted)


def break_a_number(source: Path, target: Path, column: str, wrong: str) -> None:
    """把一个数字改成文字：清洗器读不出来，提交前自检会拦住它。"""
    book = openpyxl.load_workbook(source)
    sheet = book.active
    index = [cell.value for cell in sheet[1]].index(column) + 1
    sheet.cell(row=2, column=index).value = wrong
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)


def main() -> None:
    spec = load_spec(gen.TEMPLATES / "integration.yaml")
    month_data = extend_simulation()
    if OUT.exists():
        shutil.rmtree(OUT)

    # 1 正常导入的四份（生产部含一处简称）
    write_departments(spec, month_data, OUT / "1-先导入这四份", PLANTED)
    # 2 补传更正件：同月、同部门、写法正确
    write_departments(spec, month_data, OUT / "2-补传更正", [])
    for path in (OUT / "2-补传更正").glob("*.xlsx"):
        if path.name != "生产部.xlsx":
            path.unlink()

    # 3 会被拒绝的件
    refused = OUT / "3-会被拒绝的件"
    refused.mkdir(parents=True, exist_ok=True)
    break_a_number(OUT / "1-先导入这四份" / "财务部.xlsx", refused / "财务部-金额写成文字.xlsx",
                   "本期实现收入", "约 128 万")
    shutil.copy(ROOT / "data/mock_business/monthly/2024-06-调优B/生产部.xlsx",
                refused / "生产部-这是6月的文件.xlsx")
    shutil.copy(OUT / "1-先导入这四份" / "生产部.xlsx", refused / "生产部-与已导入的完全相同.xlsx")

    (OUT / "README.md").write_text(README, encoding="utf-8")
    print(f"wrote {OUT}")


README = """# 自己上手试一遍：虚构商砼公司 2024-08

全部为**虚构数据**（一家虚构商砼公司，两个搅拌站、四个项目），由 `scripts/make_try_it_set.py` 生成，
账目互相勾稽，不对应任何真实企业。这一套不是用来评测的，是给人**手动走一遍**用的：
每一份文件都对应界面上的一个动作，以及一句应该看到的话。

## 先看这里

| 文件夹 | 里面是什么 | 拿它做什么 |
| --- | --- | --- |
| `1-先导入这四份` | 生产部 / 物资部 / 财务部 / 市场部，2024-08 | 新建批次，导入成一张跨部门总表 |
| `2-补传更正` | 生产部（写法正确的那一版） | 导入之后做「单部门补传」，看差异 |
| `3-会被拒绝的件` | 三份故意不合格的文件 | 看系统怎么拒绝、用什么话拒绝 |

## 一步一步

**① 导入**（数据 → 上传，月份选 2024-08，四份一起传）

- 提交前可以先点每个部门的「提交前自检」，四份都应通过；
- 导入后进「数据」，四行都显示「已交」，右侧「本批次的数据质量」应出现
  **跨部门不一致 · 客户名称 1 条**——生产部把「示例城建集团有限公司」写成了「示例城建集团」，
  市场部写的是全称，两边不一致时系统**两边都不采用**，那个单元格留空。

**② 看结论**（本月任务 → 发起研判 → 结论）

- 研判完成后，结论页第一层是一句话结论与四个数字；
- 关键指标带「较上期」（上期就是你之前导入过的 2024-07 批次；如果没有，会显示「无基期」而不是 0）；
- 展开任一关注项，能看到这个数字引用了哪些单元格（部门 · 文件 · 行 · 列），点进去就是原件预览。

**③ 补传更正**（数据 → 生产部那一行 → 补传更正）

- 传 `2-补传更正/生产部.xlsx`，理由写「客户单位写成了简称」；
- 应生成**新批次**：其余三个部门的来源摘要不变，只有生产部变了；
- 返回的差异里，「跨部门不一致」应由 1 条降为 0 条。

**④ 看它怎么拒绝**（`3-会被拒绝的件`）

| 文件 | 在哪用 | 应该看到 |
| --- | --- | --- |
| `财务部-金额写成文字.xlsx` | 导入前自检 | 「必须修正」里点名第 2 行的本期实现收入不是数字，并**引用清洗器自己的判定理由** |
| `生产部-这是6月的文件.xlsx` | 生产部那一行的「补传更正」 | 拒绝：这份文件报的是 2024-06，不是本批次的 2024-08——判断依据是**文件自己的行**，不是你在表单里填的月份 |
| `生产部-与已导入的完全相同.xlsx` | 同上 | 拒绝：文件与批次里那份逐字节相同，不派生新批次 |

**⑤ 看它做过什么**（记录）

- 「代理运行」把一次模型请求及其整棵调用树画成泳道：队长一条、四个部门子代理各一条；
- 「决策日志」按原话列出今天拒绝了什么、各多少次——上面那三次拒绝应当逐条出现在这里；
- 「验收评测」是生成的报告（正常路径 + 对抗路径），不是手写的。

## 口径提醒

本月依赖的「按通用做法补的口径」（增值税 13%、缺口公式、合作状态阈值、按日多行汇总）都还没有业务方确认，
所以依赖它们的数字是 **G3**。在总表视图里确认其中任意一条，依赖它的数字会升为 **G2**，**数值一个都不会变**——
确认改变的是「这句话有多硬」，不是数字本身。
"""


if __name__ == "__main__":
    main()
