"""Scale probe for #228: N production rows through the real import endpoint.

The other three departments are the retained sample; production is its rows repeated with
a unique 备注 per row (so nothing is removed as a duplicate) and dates inside 2024-07. It
imports once under the configured upload limit, once with the byte limit raised, then times
the captain's main tools on the resulting batch and prints their response sizes.

    PYTHONPATH=backend/src python scripts/scale_probe.py 199988 [csv|xlsx]
"""
import csv
import io
import json
import resource
import sys
import tempfile
import time

import openpyxl
from fastapi.testclient import TestClient

from bridgeflow.config import REPO_ROOT, settings

N = int(sys.argv[1]) if len(sys.argv) > 1 else 200_000
tmp = tempfile.mkdtemp()
settings.field_dictionary_path = str(REPO_ROOT / "data/mock_business/demo/dictionary.yaml")
settings.result_store_path = f"{tmp}/o"; settings.mapping_memory_path = f"{tmp}/m.json"; settings.column_match_path = f"{tmp}/c.json"
settings.bridgeflow_service_token = "x" * 40
for n in ("llm_provider", "llm_provider_sanitizer", "llm_provider_resolver", "llm_provider_evaluator"): setattr(settings, n, "mock")
demo = REPO_ROOT / "data/mock_business/demo"
rows = [list(r) for r in openpyxl.load_workbook(demo / "production.xlsx").active.iter_rows(values_only=True)]
header, body = rows[0], rows[1:]
buf = io.StringIO(); w = csv.writer(buf); w.writerow(header)
di = header.index("日期")
for i in range(N):
    r = list(body[i % len(body)]); r[di] = f"2024-07-{1 + i % 28:02d}"; r[header.index("备注")] = f"车次{i:06d}"; w.writerow(r)
data = buf.getvalue().encode("utf-8")
FMT = sys.argv[2] if len(sys.argv) > 2 else "csv"
if FMT == "xlsx":
    wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet()
    for row in csv.reader(io.StringIO(buf.getvalue())): ws.append(row)
    out = io.BytesIO(); wb.save(out); data = out.getvalue()
print(f"format={FMT} rows={N} bytes={len(data):,} ({len(data)/1048576:.1f} MiB) limit={settings.bridgeflow_max_upload_bytes/1048576:.0f} MiB")
from bridgeflow.api.main import app

c = TestClient(app, headers={"authorization": "Bearer " + "x" * 40})
def run(limit):
    settings.bridgeflow_max_upload_bytes = limit
    files = [("files", (f"production.{FMT}", data, "text/csv" if FMT == "csv" else "application/octet-stream"))] + [
        ("files", (f"{d}.xlsx", (demo / f"{d}.xlsx").read_bytes(), "application/octet-stream")) for d in ("procurement", "finance", "marketing")]
    t = time.perf_counter()
    r = c.post("/batches", data={"period": "2024-07", "departments": ["production", "procurement", "finance", "marketing"]}, files=files)
    return r, time.perf_counter() - t
r, dt = run(settings.bridgeflow_max_upload_bytes)
print("default limit:", r.status_code, str(r.json())[:160], f"{dt:.1f}s")
r, dt = run(200 * 1024 * 1024)
peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024
s = r.json() if r.status_code == 200 else r.text[:300]
print("raised limit:", r.status_code, f"{dt:.1f}s", f"peak_rss={peak:.0f} MiB",
      json.dumps({k: s[k] for k in ("status", "master_rows")} if isinstance(s, dict) else s, ensure_ascii=False),
      [(d["department"], d["rows"], d["quarantined"]) for d in s["departments"]] if isinstance(s, dict) else "")
if isinstance(s, dict):
    b = s["batch_id"]
    for path, body in (("/tools/batch-summary", {"batch_id": b}), ("/tools/integration-summary", {"batch_id": b}),
                       ("/tools/monthly-inbox", {"period": "2024-07"}), ("/tools/quarantine-list", {"batch_id": b}),
                       ("/tools/review-context", {"batch_id": b}),
                       ("/tools/aggregate-metric", {"batch_id": b, "metric": "production_output"})):
        t = time.perf_counter(); r = c.post(path, json=body)
        print(f"{path}: {r.status_code} {len(r.content):,} bytes {time.perf_counter()-t:.2f}s")
    print(f"final peak_rss={resource.getrusage(resource.RUSAGE_SELF).ru_maxrss/1024:.0f} MiB")
